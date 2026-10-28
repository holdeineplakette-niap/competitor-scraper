"""DEKRA source adapter.

Parses responses of the DEKRA finder API (`POST /api/locations/fromBoundaries`, a map-bounds API without key
or cookies), decides which records are in scope, exports CSV and XLSX files, and fetches the whole finder.

Scope decision (Nikos / management, 2026-10-05): include ALL German records that pass the filter below, without
validating them first; validation happens later in the competitor map. Doubtful records (untyped, thin, PSP-like,
"no vehicle testing" disclaimer) are only flagged in `doubt_flags` / `needs_review`, never dropped.

In scope = country DE and not a pure "WHG Standort" and either
  - typed "Automobil Prüfwesen und Gutachten", or
  - untyped / AT-typed ("Automobil Prüfwesen") with a vehicle-site word in the name.

Fetch plan (6 requests, `types: []` = no filter, which also returns untyped records):
  world box, Germany box, and the four quarters of the Germany box.
The world box is the data source (some DE-coded records have coordinates outside Germany's box).
The other five prove completeness: quarters == Germany box, Germany box inside world, and the world records
inside the Germany box == Germany box.

Parsing, scoping and the completeness check are pure functions. Network access happens only in DekraClient;
run(mode="simulated") uses an offline fake API. Every run writes one folder (core/paths.py,
0_instructions/storage_and_naming_convention.md).
"""

from __future__ import annotations

import csv
import json
import logging
import random
import re
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import requests
from openpyxl.utils import get_column_letter

from core import enrich, paths
from core.geo import ensure_geonames, load_plz_centroids
from core.grid import StopRun

logger = logging.getLogger(__name__)

SLUG = "dekra"
SOURCE_ORG = "DEKRA"
ID_PREFIX = "DEKRA-"
HOME_COUNTRY = "DE"  # countryResolved.code

URL = "https://www.dekra.de/api/locations/fromBoundaries"
PROFILE_HOST = "https://www.dekra.de"
HEADERS = {
    "Content-Type": "application/json",
    "Accept": "application/json",
    "Origin": "https://www.dekra.de",
    "Referer": "https://www.dekra.de/de/standorte/",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
    ),
}
PACE_S = 5.0  # between requests, plus 0-JITTER_S random
JITTER_S = 2.0
TIMEOUT_S = 60

# Location types (UUID -> name) that matter for scope; the other ~50 types are other countries / divisions.
AUTO_TYPE_ID = "2fbbb592-6bd8-4b57-860d-52b34c6cbbb1"  # Automobil Prüfwesen und Gutachten
AT_TYPE_KEY = "location_type_AT_Automotive_Testing"  # "Automobil Prüfwesen", not in DEKRA's own finder filter
WHG_TYPE_ID = "4021f9dc-8f2a-4ade-a0a4-ce74379c1089"  # WHG Standort (water-law sites)
VEHICLE_NAME_WORDS = (
    "station", "standort", "niederlassung", "außenstelle", "prüfstelle", "gutachten", "ga-büro", "automobil",
)

# Request boxes (lat/lng bounds). WORLD returns the whole universe in one response (~12 MB).
BOX_WORLD = {"latFrom": -60, "latTo": 75, "lngFrom": -180, "lngTo": 180}
BOX_GERMANY = {"latFrom": 46.9, "latTo": 55.4, "lngFrom": 5.5, "lngTo": 15.4}
_MID_LAT = (BOX_GERMANY["latFrom"] + BOX_GERMANY["latTo"]) / 2
_MID_LNG = (BOX_GERMANY["lngFrom"] + BOX_GERMANY["lngTo"]) / 2
QUARTERS = {  # q1 north-west, q2 north-east, q3 south-west, q4 south-east
    "q1": {"latFrom": _MID_LAT, "latTo": BOX_GERMANY["latTo"], "lngFrom": BOX_GERMANY["lngFrom"], "lngTo": _MID_LNG},
    "q2": {"latFrom": _MID_LAT, "latTo": BOX_GERMANY["latTo"], "lngFrom": _MID_LNG, "lngTo": BOX_GERMANY["lngTo"]},
    "q3": {"latFrom": BOX_GERMANY["latFrom"], "latTo": _MID_LAT, "lngFrom": BOX_GERMANY["lngFrom"], "lngTo": _MID_LNG},
    "q4": {"latFrom": BOX_GERMANY["latFrom"], "latTo": _MID_LAT, "lngFrom": _MID_LNG, "lngTo": BOX_GERMANY["lngTo"]},
}
FULL_REQUESTS = {"world": BOX_WORLD, "de_box": BOX_GERMANY, **QUARTERS}
SAMPLE_REQUESTS = {"q4": QUARTERS["q4"]}  # the smallest quarter, ~1 MB
MODES = ("simulated", "sample", "full")
FULL_RUN_NOT_CONFIRMED = "A full DEKRA run makes 6 live requests (~30 MB). Pass confirm_full_run=True to start it."

MAX_GERMAN_DROP = 0.10  # alert if in-scope rows drop more than this vs the last passing full run

SITE_TYPES = ("pruefstelle", "partner_office", "unknown")
FLAG_UNTYPED = "untyped"
FLAG_AT_TYPE = "at_type"
FLAG_THIN = "thin"
FLAG_PSP = "psp_suspect"
FLAG_DISCLAIMER = "no_vehicle_disclaimer"
FLAG_NON_DEKRA_OPERATOR = "non_dekra_operator"
# Flags that also set needs_review. non_dekra_operator (partner offices) is informational only.
REVIEW_FLAGS = (FLAG_UNTYPED, FLAG_AT_TYPE, FLAG_THIN, FLAG_PSP, FLAG_DISCLAIMER)

COLUMNS = [
    "location_id",
    "source_org",
    "source_id",
    "name",
    "name_addon",
    "name_addon2",
    "street",
    "zip",
    "city",
    "state",
    "country",
    "lat",
    "lng",
    "phone",
    "mobile",
    "email",
    "website",
    "dekra_profile_url",
    "booking_url",
    "opening_hours",
    "services",
    "dekra_types",
    "nfz",
    "hu_mention",
    "same_address_count",
    "doubt_flags",
    "site_type",
    "site_type_evidence",
    "needs_review",
    "needs_review_reason",
    "in_scope",
    "scope_reason",
    "scraped_at",
]

REASON_OUTSIDE_GERMANY = "outside Germany"
REASON_COUNTRY_MISSING = "country missing"
REASON_WHG_ONLY = "WHG site only"
REASON_NOT_VEHICLE_NAME = "untyped or AT-typed, no vehicle-site name"
REASON_OTHER_TYPE = "type not vehicle-related"

_HU_WORD = re.compile(r"\bHU\b")
_HU_LONG = re.compile(r"hauptuntersuchung", re.I)
_HU_FIELDS = ("freeText1", "freeText2", "services", "currentInformation", "openingHoursAreas", "name", "name2")
_PSP_PATTERN = re.compile(
    r"\(|\b(Autohaus|Autozentrum|Garage|Werkstatt|Kfz-Service|KIA|Toyota|Opel|Ford|Mercedes|VW|Volkswagen|Renault|"
    r"Hyundai|BMW|Audi|Skoda|Seat|Honda|Mazda|Nissan|Suzuki|Volvo|Peugeot|Fiat|Gel[aä]nde)\b",
    re.I,
)
_DISCLAIMER = re.compile(r"keine (Kfz-Pr[uü]fstelle|Fahrzeugpr[uü]fung)|nur Industrie", re.I)
_WEEKDAYS = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
_WEEKDAY_SHORT = {"monday": "Mo", "tuesday": "Di", "wednesday": "Mi", "thursday": "Do", "friday": "Fr",
                  "saturday": "Sa", "sunday": "So"}


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def _clean(value) -> str | None:
    """Strip strings and collapse inner whitespace; empty strings and None become None."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _is_empty(v) -> bool:
    if v is None or v == "" or v == [] or v == {}:
        return True
    if isinstance(v, dict):
        kind = v.get("type")
        if kind == "RichText":
            return not v.get("elements")
        if kind == "Date":
            return not v.get("date")
        if kind == "Link":
            return not v.get("to")
    return False


def _strings(v):
    if isinstance(v, str):
        yield v
    elif isinstance(v, dict):
        for x in v.values():
            yield from _strings(x)
    elif isinstance(v, list):
        for x in v:
            yield from _strings(x)


def _rich_text(v) -> str:
    """Flatten a CMS RichText value to plain text (text leaves only)."""
    out: list[str] = []

    def rec(x):
        if isinstance(x, str):
            out.append(x)
        elif isinstance(x, dict):
            for key in ("elements", "content"):
                if key in x:
                    rec(x[key])
        elif isinstance(x, list):
            for y in x:
                rec(y)

    rec(v)
    return re.sub(r"\s+", " ", " ".join(out)).strip()


def _link(value) -> str | None:
    """Full URL of a CMS Link object, or None."""
    if not isinstance(value, dict):
        return None
    to = _clean(value.get("to"))
    if not to:
        return None
    return PROFILE_HOST + to if to.startswith("/") else to


def _country(rec: dict) -> str | None:
    return (rec.get("countryResolved") or {}).get("code")


def _types(rec: dict) -> list[dict]:
    return rec.get("locationTypesResolved") or []


def _street(rec: dict) -> str | None:
    street = _clean(rec.get("street"))
    number = _clean(rec.get("houseNumber"))
    if street and number and number not in street:
        return f"{street} {number}"
    return street or number


def _zip(rec: dict) -> str | None:
    """Trimmed PLZ; German PLZ with a lost leading zero (4 digits) are padded to 5."""
    z = _clean(rec.get("zipCode"))
    if z and _country(rec) == HOME_COUNTRY and z.isdigit() and len(z) < 5:
        z = z.zfill(5)
    return z


def opening_hours(rec: dict) -> str | None:
    """Compact text from the structured opening-hours sections, e.g. 'Fahrzeugprüfungen: Mo–Do 08:00–18:00, Fr 08:00–16:00'.

    The times in the source are stored as 1970-01-01THH:MM:00.000Z but are local times, so no time-zone shift.
    """
    parts = []
    for area in rec.get("openingHoursAreas") or []:
        data = area.get("data") or {}
        by_hours: dict[str, set[str]] = {}
        for hours in data.get("openingHours") or []:
            h = hours.get("data") or {}
            start = ((h.get("from") or {}).get("date") or "")[11:16]
            end = ((h.get("to") or {}).get("date") or "")[11:16]
            if not start or not end:
                continue
            by_hours.setdefault(f"{start}–{end}", set()).update(o.get("key") for o in h.get("weekdays") or [])
        spans = []
        for span, days in by_hours.items():
            ordered = [d for d in _WEEKDAYS if d in days]
            groups, current = [], []
            for d in ordered:
                if current and _WEEKDAYS.index(d) != _WEEKDAYS.index(current[-1]) + 1:
                    groups.append(current)
                    current = []
                current.append(d)
            if current:
                groups.append(current)
            for g in groups:
                label = _WEEKDAY_SHORT[g[0]] if len(g) == 1 else f"{_WEEKDAY_SHORT[g[0]]}–{_WEEKDAY_SHORT[g[-1]]}"
                spans.append((_WEEKDAYS.index(g[0]), f"{label} {span}"))
        if spans:
            text = ", ".join(s for _, s in sorted(spans))
            headline = _clean(data.get("headline"))
            parts.append(f"{headline}: {text}" if headline else text)
    return " | ".join(parts) or None


def hu_mention(rec: dict) -> str | None:
    """Names of the text fields that mention 'Hauptuntersuchung' or the word 'HU'."""
    hits = []
    for field in _HU_FIELDS:
        text = " ".join(_strings(rec.get(field)))
        if _HU_LONG.search(text) or _HU_WORD.search(text):
            hits.append(field)
    return "; ".join(hits) or None


# --------------------------------------------------------------------------
# Scope and flags
# --------------------------------------------------------------------------

def vehicle_name_word(name: str | None) -> str | None:
    low = (name or "").lower()
    return next((w for w in VEHICLE_NAME_WORDS if w in low), None)


def classify(rec: dict) -> tuple[bool, str, str, str]:
    """(in_scope, scope_reason, site_type, evidence). See the module docstring for the scope rule."""
    country = _country(rec)
    if not country:
        return False, REASON_COUNTRY_MISSING, "unknown", ""
    if country != HOME_COUNTRY:
        return False, REASON_OUTSIDE_GERMANY, "unknown", ""
    types = _types(rec)
    ids = {t.get("id") for t in types}
    if ids == {WHG_TYPE_ID}:
        return False, REASON_WHG_ONLY, "unknown", ""
    if AUTO_TYPE_ID in ids:
        return True, "", "pruefstelle", "dekra_type:automotive_testing"
    at_typed = any(t.get("key") == AT_TYPE_KEY for t in types)
    if not types or at_typed:
        word = vehicle_name_word(rec.get("name"))
        kind = "at_type" if at_typed else "untyped"
        if word:
            return True, "", "unknown", f"{kind}; name:{word}"
        return False, REASON_NOT_VEHICLE_NAME, "unknown", ""
    names = ", ".join(t.get("name", "?") for t in types)
    return False, f"{REASON_OTHER_TYPE} ({names})", "unknown", ""


def flags_of(rec: dict) -> list[str]:
    """Cheap doubt markers for the later validation. They never exclude a record."""
    flags = []
    types = _types(rec)
    ids = {t.get("id") for t in types}
    if not types:
        flags.append(FLAG_UNTYPED)
    elif AUTO_TYPE_ID not in ids and any(t.get("key") == AT_TYPE_KEY for t in types):
        flags.append(FLAG_AT_TYPE)
    has_hours = bool(opening_hours(rec))
    if not _link(rec.get("linkWebsite")) and not _link(rec.get("linkOnlineAppointment")) and not has_hours \
            and _is_empty(rec.get("freeText1")):
        flags.append(FLAG_THIN)
    if _PSP_PATTERN.search(f"{rec.get('name') or ''} {rec.get('name2') or ''}"):
        flags.append(FLAG_PSP)
    text = f"{_rich_text(rec.get('freeText1'))} {_rich_text(rec.get('currentInformation'))}"
    if _DISCLAIMER.search(text):
        flags.append(FLAG_DISCLAIMER)
    operator = _clean(rec.get("name2"))
    if operator and "dekra" not in operator.lower():
        flags.append(FLAG_NON_DEKRA_OPERATOR)
    return flags


# --------------------------------------------------------------------------
# Parsing
# --------------------------------------------------------------------------

def parse_record(rec: dict, scraped_at: str) -> dict:
    """One API record -> one row with the columns in COLUMNS. Contact fields are kept (personal data!)."""
    in_scope, scope_reason, site_type, evidence = classify(rec)
    flags = flags_of(rec)
    review = [f for f in flags if f in REVIEW_FLAGS]
    source_id = _clean(rec.get("id"))
    return {
        "location_id": f"{ID_PREFIX}{source_id}",
        "source_org": SOURCE_ORG,
        "source_id": source_id,
        "name": _clean(rec.get("name")),
        "name_addon": _clean(rec.get("name2")),
        "name_addon2": _clean(rec.get("additionalAddress")),
        "street": _street(rec),
        "zip": _zip(rec),
        "city": _clean(rec.get("city")),
        "state": None,  # filled from the PLZ (core.enrich); DEKRA's own `state` field is 4% filled
        "country": _country(rec),
        "lat": rec.get("latitude"),
        "lng": rec.get("longitude"),
        "phone": _clean(rec.get("phone1")),
        "mobile": _clean(rec.get("phoneMobile")),
        "email": _clean(rec.get("email")),
        "website": None,  # DEKRA lists no own website per location, only its profile page
        "dekra_profile_url": _link(rec.get("linkWebsite")),
        "booking_url": _link(rec.get("linkOnlineAppointment")),
        "opening_hours": opening_hours(rec),
        "services": "; ".join(o.get("value", "") for o in rec.get("specialServicesList") or []) or None,
        "dekra_types": " | ".join(t.get("name", "") for t in _types(rec)) or None,
        "nfz": bool(rec.get("suitableForCommercialVehicles")),
        "hu_mention": hu_mention(rec),
        "same_address_count": None,  # set by add_address_counts
        "doubt_flags": "; ".join(flags) or None,
        "site_type": site_type,
        "site_type_evidence": evidence or None,
        "needs_review": bool(review),
        "needs_review_reason": "; ".join(review) or None,
        "in_scope": in_scope,
        "scope_reason": scope_reason or None,
        "scraped_at": scraped_at,
    }


def _address_key(row) -> tuple[str, str]:
    # missing values arrive as None or NaN depending on the column's dtype
    street = (row["street"] if isinstance(row["street"], str) else "").lower().replace("ß", "ss")
    street = re.sub(r"stra(ss|ß)e|str\.", "str", street)
    zip_code = row["zip"] if isinstance(row["zip"], str) else ""
    return re.sub(r"[^a-z0-9äöü]", "", street), zip_code.replace(" ", "")


def add_address_counts(df: pd.DataFrame) -> pd.DataFrame:
    """same_address_count = rows (all countries, in or out of scope) with the same normalised street + PLZ, itself included."""
    out = df.copy()
    keys = [_address_key(r) for _, r in out.iterrows()]
    counts: dict = {}
    for k in keys:
        counts[k] = counts.get(k, 0) + 1
    # a record without street or PLZ has no address to share, so it counts as 1
    out["same_address_count"] = [counts[k] if k[0] and k[1] else 1 for k in keys]
    return out


def parse_records(records: list[dict], scraped_at: str) -> pd.DataFrame:
    df = pd.DataFrame([parse_record(r, scraped_at) for r in records], columns=COLUMNS)
    return add_address_counts(df) if len(df) else df


def load_response(path: str | Path) -> dict:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def enrich_german(df: pd.DataFrame, centroids: pd.DataFrame) -> pd.DataFrame:
    """core.enrich for German rows only (state, PLZ distance, PLZ flags)."""
    return enrich.enrich(df, centroids, check=df["country"].eq(HOME_COUNTRY))


# --------------------------------------------------------------------------
# Completeness
# --------------------------------------------------------------------------

def in_box(rec: dict, box: dict) -> bool:
    lat, lng = rec.get("latitude"), rec.get("longitude")
    if lat is None or lng is None:
        return False
    return box["latFrom"] <= lat <= box["latTo"] and box["lngFrom"] <= lng <= box["lngTo"]


def check_completeness(responses: dict[str, list[dict]]) -> dict:
    """Cross-check the six responses by record id (see the module docstring)."""
    ids = {name: {r["id"] for r in recs} for name, recs in responses.items()}
    duplicates = {name: len(recs) - len(ids[name]) for name, recs in responses.items() if len(recs) != len(ids[name])}
    quarters = set().union(*(ids[q] for q in QUARTERS))
    world_in_box = {r["id"] for r in responses["world"] if in_box(r, BOX_GERMANY)}
    checks = {
        "no_duplicate_ids": not duplicates,
        "quarters_equal_germany_box": quarters == ids["de_box"],
        "germany_box_inside_world": ids["de_box"] <= ids["world"],
        "world_in_box_equals_germany_box": world_in_box == ids["de_box"],
    }
    return {
        "passed": all(checks.values()),
        "checks": checks,
        "counts": {name: len(v) for name, v in ids.items()},
        "duplicates": duplicates,
        "quarters_only": len(quarters - ids["de_box"]),
        "box_only_vs_quarters": len(ids["de_box"] - quarters),
        "box_not_in_world": len(ids["de_box"] - ids["world"]),
        "world_in_box_not_in_box": len(world_in_box - ids["de_box"]),
        "box_not_in_world_in_box": len(ids["de_box"] - world_in_box),
    }


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------

XLSX_TEXT_COLUMNS = ["source_id", "location_id", "zip"]  # stored as text so Excel keeps leading zeros
XLSX_MAX_COL_WIDTH = 60


def safe_path(path: str | Path, renamed: list | None = None) -> Path:
    """Never overwrite: if path exists, return name_2, name_3, ... and log the rename."""
    path = Path(path)
    if not path.exists():
        return path
    n = 2
    while (candidate := path.with_name(f"{path.stem}_{n}{path.suffix}")).exists():
        n += 1
    logger.warning("%s already exists, writing %s instead", path, candidate.name)
    if renamed is not None:
        renamed.append((str(path), str(candidate)))
    return candidate


def _write_xlsx(frame: pd.DataFrame, path: Path, sheet_name: str) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(frame.columns, start=1):
            letter = get_column_letter(idx)
            if col in XLSX_TEXT_COLUMNS:
                for cell in ws[letter][1:]:
                    cell.number_format = "@"
                    if cell.value is not None:
                        cell.value = str(cell.value)
            width = max([len(str(col))] + [len(str(v)) for v in frame[col].dropna()])
            ws.column_dimensions[letter].width = min(width + 2, XLSX_MAX_COL_WIDTH)


EXPORT_KEYS = ("all_csv", "all_xlsx", "pruefstellen_csv", "pruefstellen_xlsx")


def export(
    df: pd.DataFrame,
    out_dir: str | Path,
    run_date: str | None = None,
    renamed: list | None = None,
    tag: str | None = None,
    only: tuple[str, ...] | list[str] | None = None,
) -> dict:
    """Write all rows (every record the source lists) and in-scope rows to CSV (utf-8-sig, ';') and XLSX.

    only: a subset of EXPORT_KEYS to write (default: all four), e.g. ("pruefstellen_csv",).
    Existing files are never overwritten. tag, if given, is appended to the file names. Returns {name: (path, rows)}.
    """
    wanted = set(EXPORT_KEYS if only is None else only)
    unknown = wanted - set(EXPORT_KEYS)
    if unknown:
        raise ValueError(f"unknown export key(s) {sorted(unknown)}; choose from {EXPORT_KEYS}")
    run_date = run_date or date.today().isoformat()
    suffix = f"_{tag}" if tag else ""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    result = {}
    for key, frame in {"all": df, "pruefstellen": df[df["in_scope"]]}.items():
        if f"{key}_csv" in wanted:
            csv_path = safe_path(out_dir / f"{SLUG}_{key}_{run_date}{suffix}.csv", renamed)
            frame.to_csv(csv_path, sep=";", encoding="utf-8-sig", index=False)
            result[f"{key}_csv"] = (csv_path, len(frame))
        if f"{key}_xlsx" in wanted:
            xlsx_path = safe_path(out_dir / f"{SLUG}_{key}_{run_date}{suffix}.xlsx", renamed)
            _write_xlsx(frame, xlsx_path, sheet_name=key)
            result[f"{key}_xlsx"] = (xlsx_path, len(frame))
    return result


# --------------------------------------------------------------------------
# Fetch
# --------------------------------------------------------------------------

class ApiError(StopRun):
    """HTTP error, timeout, non-JSON answer or unexpected response shape."""


def build_body(boundaries: dict) -> dict:
    """Request body: no type filter (so untyped records come back too), no Nfz restriction."""
    return {
        "locale": "de_DE",
        "searchedLocation": {"lat": 52.530741, "lng": 13.4058114},
        "tagFilters": [],
        "initialSearch": False,
        "boundaries": dict(boundaries),
        "types": [],
        "commercialVehicles": False,
        "locationsWithoutType": False,
    }


def check_shape(data) -> list[dict]:
    """The record list of a response; stops the run on an unexpected shape or isNextLocation == true."""
    if not isinstance(data, dict) or not isinstance(data.get("locations"), list) or "isNextLocation" not in data:
        raise ApiError(f"unexpected response shape: {str(data)[:200]!r}")
    if data["isNextLocation"]:
        raise ApiError("isNextLocation is true: the response is incomplete, the adapter does not page")
    return data["locations"]


class DekraClient:
    """Sequential POST client: PACE_S + jitter between requests, one retry on timeout/connection error."""

    def __init__(self, pace_s: float = PACE_S, jitter_s: float = JITTER_S, sleep=time.sleep):
        self.pace_s, self.jitter_s, self._sleep = pace_s, jitter_s, sleep
        self._first = True

    def post(self, body: dict) -> tuple[dict, dict]:
        if not self._first:
            self._sleep(self.pace_s + random.uniform(0, self.jitter_s))
        self._first = False
        t0 = time.time()
        for attempt in (1, 2):
            try:
                resp = requests.post(URL, headers=HEADERS, data=json.dumps(body), timeout=TIMEOUT_S)
                break
            except (requests.Timeout, requests.ConnectionError) as exc:
                if attempt == 2:
                    raise ApiError(f"DEKRA API unreachable after 2 attempts: {type(exc).__name__}") from exc
                self._sleep(self.pace_s + random.uniform(0, self.jitter_s))
        if resp.status_code >= 400:
            raise ApiError(f"HTTP {resp.status_code}: {resp.text[:300]!r}")
        try:
            data = resp.json()
        except ValueError as exc:
            raise ApiError(f"HTTP {resp.status_code} but the body is not JSON: {resp.text[:200]!r}") from exc
        return data, {"http_status": resp.status_code, "bytes": len(resp.content), "elapsed_s": round(time.time() - t0, 1)}


def _fake_record(i: int, rng: random.Random) -> dict:
    """A synthetic record in the API's shape. No real data."""
    kind = i % 10
    in_germany = kind != 9
    lat = rng.uniform(47.3, 54.9) if in_germany else rng.uniform(46.0, 47.0)
    lng = rng.uniform(6.0, 14.8) if in_germany else rng.uniform(17.0, 20.0)
    code = "DE" if in_germany else "SK"
    if kind in (0, 1, 2, 3, 4, 5):
        types = [{"id": AUTO_TYPE_ID, "key": "automotive_testing", "name": "Automobil Prüfwesen und Gutachten"}]
        name = f"Station Teststadt {i}"
    elif kind == 6:
        types, name = [], f"Standort Testdorf {i}"  # untyped, vehicle-like name -> in scope, flagged
    elif kind == 7:
        types, name = [{"id": WHG_TYPE_ID, "key": "whg_location", "name": "WHG Standort"}], f"Anlagenbau {i} GmbH"
    elif kind == 8:
        types, name = [], f"DEKRA Akademie {i}"  # untyped, no vehicle word -> out of scope
    else:
        types, name = [], f"DEKRA Slovensko {i}"
    return {
        "id": f"00000000-0000-4000-8000-{i:012d}",
        "name": name,
        "name2": "DEKRA Automobil GmbH",
        "street": "Teststr.",
        "houseNumber": str(i),
        "zipCode": f"{rng.randint(1000, 99999):05d}",
        "city": "Teststadt",
        "latitude": lat,
        "longitude": lng,
        "countryResolved": {"code": code, "name": "Test"},
        "locationTypesResolved": types,
        "suitableForCommercialVehicles": i % 3 == 0,
        "specialServicesList": [],
        "openingHoursAreas": [],
        "freeText1": {"type": "RichText", "elements": []},
        "currentInformation": {"type": "RichText", "elements": []},
        "linkWebsite": {"type": "Link", "to": f"https://www.dekra.de/de/test-{i}"} if i % 4 else None,
        "linkOnlineAppointment": None,
    }


def simulated_records(n: int = 100, seed: int = 1) -> list[dict]:
    rng = random.Random(seed)
    return [_fake_record(i, rng) for i in range(n)]


class FakeDekraApi:
    """Offline stand-in for DekraClient: answers every request from a fixed record list (the map-bounds filter)."""

    def __init__(self, records: list[dict]):
        self.records = records
        self.calls: list[dict] = []

    def post(self, body: dict) -> tuple[dict, dict]:
        self.calls.append(body)
        box = body["boundaries"]
        data = {"isNextLocation": False, "locations": [r for r in self.records if in_box(r, box)]}
        return data, {"http_status": 200, "bytes": len(json.dumps(data)), "elapsed_s": 0.0}


# --------------------------------------------------------------------------
# Run
# --------------------------------------------------------------------------

QUERY_LOG_COLUMNS = ["seq", "request", "time", "http_status", "bytes", "records", "elapsed_s", "status"]


@dataclass
class RunResult:
    run_folder: Path
    summary: dict
    frame: pd.DataFrame
    files: dict


def _relative(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError:
        return str(path)


def _baseline(output_dir: Path, exclude: Path | None) -> dict | None:
    """The newest earlier official run whose run_summary.json says PASS."""
    if not output_dir.is_dir():
        return None
    for folder in sorted(output_dir.iterdir(), reverse=True):
        if folder == exclude or not (folder / paths.RUN_SUMMARY_JSON).exists():
            continue
        try:
            summary = json.loads((folder / paths.RUN_SUMMARY_JSON).read_text(encoding="utf-8"))
        except ValueError:
            continue
        if summary.get("mode") == "full" and summary.get("status") == "PASS" and summary.get("in_scope") is not None:
            return {"run_folder": folder.name, "run_date": summary.get("run_date"), "in_scope": summary["in_scope"]}
    return None


def _counts(series: pd.Series) -> dict:
    return {str(k): int(v) for k, v in series.value_counts(dropna=True).items()}


def _flag_counts(df: pd.DataFrame) -> dict:
    counts: dict[str, int] = {}
    for value in df.loc[df["in_scope"], "doubt_flags"].dropna():
        for flag in value.split("; "):
            counts[flag] = counts.get(flag, 0) + 1
    return counts


def format_run_summary(summary: dict) -> str:
    lines = [f"DEKRA run summary ({summary['mode']}): {summary['outcome']}", "=" * 50,
             f"Run date:           {summary['run_date']}  (started {summary['started_at']}, {summary['runtime_s']} s)",
             f"Requests made:      {summary['requests_made']}  ({', '.join(f'{k} {v}' for k, v in summary['records_per_request'].items())})"]
    if summary["status_reasons"]:
        lines += ["FAIL because:"] + [f"  ! {r}" for r in summary["status_reasons"]]
    comp = summary.get("completeness")
    if comp:
        lines += ["", "Completeness (by record id):"] + [f"  {'ok  ' if ok else 'FAIL'} {name}" for name, ok in comp["checks"].items()]
    lines += ["",
              f"Collected:          {summary['rows_total']} unique records, {summary['rows_germany']} with country DE",
              f"In scope:           {summary['in_scope']} (DE, not WHG-only, Automobil-typed or untyped/AT with a vehicle-site name)",
              f"Site types (in scope): {summary['site_types']}",
              f"Flags (in scope):   {summary['flags']}",
              f"Out of scope:       {summary['scope_reasons']}",
              f"Needs review:       {summary['needs_review_in_scope']} in-scope rows (flags only; nothing was excluded)"]
    if summary.get("outside_box"):
        lines.append(f"DE-coded in-scope rows outside Germany's box: {summary['outside_box']} (only the world request returns them)")
    lines.append(f"Drop check:         {summary['drop_check']}")
    if summary.get("baseline"):
        lines.append(f"Baseline:           {summary['baseline']['run_folder']} ({summary['baseline']['in_scope']} in scope)")
    if not summary.get("save_raw") and summary["mode"] != "simulated":
        lines.append("Raw responses:      not saved (save_raw=False); this run cannot be re-parsed offline")
    lines += ["", "Files:"] + [f"  {k:<18}{v}" for k, v in summary["files"].items()]
    return "\n".join(lines)


def run(
    mode: str = "simulated",
    *,
    confirm_full_run: bool = False,
    root: str | Path | None = None,
    started_at: datetime | None = None,
    api=None,
    records: list[dict] | None = None,
    centroids: pd.DataFrame | None = None,
    verbose: bool = True,
    exports: tuple[str, ...] | list[str] = EXPORT_KEYS,
    save_raw: bool | None = None,
) -> RunResult:
    """One run = one folder. simulated: offline fake API (no raw.jsonl). sample: 1 live request (quarter q4).
    full: 6 live requests and the completeness check; needs confirm_full_run=True.

    exports: which export files to write (subset of EXPORT_KEYS; default all four).
    save_raw: keep the raw responses in raw.jsonl (default: yes for live runs, no for simulated; a full run's
    raw.jsonl is ~30 MB). Without it nothing can be re-parsed offline later. The checks always use the full
    responses in memory, whatever is saved. query_log.csv and run_summary.json/.md are always written.
    """
    if mode not in MODES:
        raise ValueError(f"mode must be one of {MODES}, got {mode!r}")
    if mode == "full" and not confirm_full_run:
        raise RuntimeError(FULL_RUN_NOT_CONFIRMED)
    exports = tuple(exports)
    if set(exports) - set(EXPORT_KEYS):  # checked first: no run folder is created for a bad setting
        raise ValueError(f"unknown export key(s) {sorted(set(exports) - set(EXPORT_KEYS))}; choose from {EXPORT_KEYS}")
    root = Path(root) if root is not None else paths.project_root()
    started = started_at or datetime.now()
    t0 = time.time()
    scraped_at = started.isoformat(timespec="seconds")
    run_date = started.date().isoformat()
    run_folder = paths.new_run_dir(SLUG, mode, started, root)

    if mode == "simulated":
        api = api or FakeDekraApi(records if records is not None else simulated_records())
    else:
        api = api or DekraClient()
    plan = SAMPLE_REQUESTS if mode == "sample" else FULL_REQUESTS
    keep_raw = (mode != "simulated") if save_raw is None else bool(save_raw)

    def say(msg):
        if verbose:
            print(msg, flush=True)

    responses: dict[str, list[dict]] = {}
    log_rows: list[dict] = []
    abort_reason = None
    raw_file = open(run_folder / paths.RAW_FILE, "x", encoding="utf-8") if keep_raw else None
    try:
        for seq, (name, box) in enumerate(plan.items(), start=1):
            say(f"[{seq}/{len(plan)}] {name}")
            row = {"seq": seq, "request": name, "time": datetime.now().isoformat(timespec="seconds")}
            try:
                data, info = api.post(build_body(box))
                recs = check_shape(data)
            except StopRun as exc:
                log_rows.append({**row, "status": "error"})
                abort_reason = f"{name}: {exc}"
                break
            if raw_file:
                raw_file.write(json.dumps({"file": f"{name}.json", "response": data}, ensure_ascii=False) + "\n")
            responses[name] = recs
            log_rows.append({**row, **info, "records": len(recs), "status": "ok"})
            say(f"    {len(recs)} records, {info['bytes'] / 1e6:.1f} MB")
    finally:
        if raw_file:
            raw_file.close()
    with open(run_folder / paths.QUERY_LOG_FILE, "x", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=QUERY_LOG_COLUMNS)
        writer.writeheader()
        writer.writerows(log_rows)

    # ---- parse, enrich
    source = responses.get("world") or [r for recs in responses.values() for r in recs]
    unique = list({r["id"]: r for r in source}.values())
    df = parse_records(unique, scraped_at)
    if len(df):
        if centroids is None:
            centroids = load_plz_centroids(ensure_geonames(paths.geonames_dir(root)))
        df = enrich_german(df, centroids)

    # ---- status
    completeness = check_completeness(responses) if mode != "sample" and not abort_reason else None
    reasons = []
    if abort_reason:
        reasons.append(f"run aborted: {abort_reason}")
    if completeness and not completeness["passed"]:
        reasons += [f"completeness check failed: {k}" for k, ok in completeness["checks"].items() if not ok]
    if mode == "sample":
        status, outcome = "SAMPLE", "SAMPLE (one quarter only, no completeness proof)"
    elif reasons:
        status, outcome = "FAIL", "FAIL"
    else:
        status, outcome = "PASS", "PASS (6 responses consistent)"

    tag = None if status in ("PASS", "SAMPLE") else "unverified"
    renamed: list = []
    written = export(df, run_folder, run_date, renamed, tag, only=exports) if len(df) else {}

    in_scope_df = df[df["in_scope"]] if len(df) else df
    german = df[df["country"].eq(HOME_COUNTRY)] if len(df) else df
    outside_box = 0
    if "world" in responses and len(df):
        by_id = {r["id"]: r for r in unique}
        outside_box = sum(1 for loc in in_scope_df["source_id"] if not in_box(by_id[loc], BOX_GERMANY))
    output_dir = paths.output_dir(SLUG, "full", root)
    baseline = _baseline(output_dir, run_folder) if mode == "full" else None
    drop_check, drop_alert = "not run", False
    if mode in ("full", "simulated") and status != "FAIL":
        if baseline:
            change = (len(in_scope_df) - baseline["in_scope"]) / baseline["in_scope"] if baseline["in_scope"] else 0.0
            drop_alert = change < -MAX_GERMAN_DROP
            drop_check = f"{'ALERT' if drop_alert else 'ok'}: {baseline['in_scope']} -> {len(in_scope_df)} in scope ({change:+.1%}, alert below -{MAX_GERMAN_DROP:.0%})"
        else:
            drop_check = "no baseline (no earlier passing full run)"
    elif status == "FAIL":
        drop_check = "not run (run failed)"

    summary = {
        "source": SOURCE_ORG,
        "mode": mode,
        "status": status,
        "outcome": outcome,
        "status_reasons": reasons,
        "aborted": bool(abort_reason),
        "run_date": run_date,
        "started_at": scraped_at,
        "runtime_s": round(time.time() - t0, 1),
        "requests_made": len(log_rows),
        "records_per_request": {r["request"]: r.get("records") for r in log_rows if r["status"] == "ok"},
        "completeness": completeness,
        "rows_total": int(len(df)),
        "rows_germany": int(len(german)),
        "in_scope": int(len(in_scope_df)),
        "site_types": _counts(in_scope_df["site_type"]) if len(df) else {},
        "flags": _flag_counts(df) if len(df) else {},
        "scope_reasons": _counts(df.loc[~df["in_scope"], "scope_reason"].map(lambda v: str(v).split(" (")[0])) if len(df) else {},
        "needs_review_in_scope": int(in_scope_df["needs_review"].sum()) if len(df) else 0,
        "review_reasons": enrich.reason_counts(in_scope_df) if len(df) else {},
        "outside_box": outside_box,
        "drop_check": drop_check,
        "german_drop_alert": drop_alert,
        "baseline": baseline,
        "save_raw": keep_raw,
        "exports": list(exports),
        "files": {"run_folder": _relative(run_folder, root)}
        | {key: _relative(path, root) for key, (path, _rows) in written.items()},
    }
    text = format_run_summary(summary)
    (run_folder / paths.RUN_SUMMARY_JSON).write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (run_folder / paths.RUN_SUMMARY_MD).write_text(f"# Run summary: {run_folder.name}\n\n```text\n{text}\n```\n", encoding="utf-8")
    say(text)
    return RunResult(run_folder=run_folder, summary=summary, frame=df, files=written)


def write_recon_run(
    df: pd.DataFrame,
    source_file: str | Path,
    root: str | Path | None = None,
    run_date: str | None = None,
    started_at: datetime | None = None,
    exports: tuple[str, ...] | list[str] = EXPORT_KEYS,
) -> tuple[Path, dict]:
    """Mode "recon": export an offline-parsed saved response into a run folder of its own in test_output/.

    No requests. Writes the chosen exports (default all four) plus run_summary.json/.md; returns (run folder, export result).
    """
    root = Path(root) if root is not None else paths.project_root()
    started = started_at or datetime.now()
    run_date = run_date or started.date().isoformat()
    run_folder = paths.new_run_dir(SLUG, "recon", started, root)
    written = export(df, run_folder, run_date, only=exports)
    in_scope = df[df["in_scope"]]
    summary = {
        "source": SOURCE_ORG,
        "mode": "recon",
        "run_date": run_date,
        "started_at": started.isoformat(timespec="seconds"),
        "source_file": _relative(Path(source_file), root),
        "rows": int(len(df)),
        "in_scope": int(len(in_scope)),
        "site_types": _counts(in_scope["site_type"]),
        "flags": _flag_counts(df),
        "files": {"run_folder": _relative(run_folder, root)}
        | {key: _relative(path, root) for key, (path, _rows) in written.items()},
    }
    (run_folder / paths.RUN_SUMMARY_JSON).write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    text = "\n".join(
        [f"DEKRA recon export (no network): {summary['rows']} rows, {summary['in_scope']} in scope",
         f"Source file: {summary['source_file']}", f"Site types (in scope): {summary['site_types']}",
         f"Flags (in scope): {summary['flags']}", "", "Files:"]
        + [f"  {k:<18}{v}" for k, v in summary["files"].items()]
    )
    (run_folder / paths.RUN_SUMMARY_MD).write_text(f"# Run summary: {run_folder.name}\n\n```text\n{text}\n```\n", encoding="utf-8")
    return run_folder, written
