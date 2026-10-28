"""Offline: draw the DEKRA labelling sample from the saved recon universe (no network requests).

    .venv\\Scripts\\python.exe 1_scrapers/dekra/src/labelling_sample.py [recon_run_folder_name]

Reads requests/r09_response.json of the recon run (via recon.load_responses), writes into 4_reference/ground_truth/:
    dekra_labelling_sample_<date>.xlsx / .csv      the sample (sheet "Sample" + sheet "Guide")
    dekra_labelling_s2_matches_<date>.csv          every record the person-name heuristic (S2) matched
A file is never overwritten: if a name exists, all outputs of the run get the same _2/_3 suffix.
The files contain personal data (person-named engineering offices): they stay in git-ignored 4_reference/.
"""

from __future__ import annotations

import csv
import random
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from urllib.parse import quote_plus

sys.path.insert(0, str(Path(__file__).resolve().parent))

import recon  # noqa: E402  (parsing helpers: load_responses, is_empty, strings)
from core import paths  # noqa: E402
from core.geo import haversine_km, load_plz_centroids  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

SEED = 42
DEFAULT_RUN = "2026-09-30_1831_recon"
AUTO_TYPE_ID = "2fbbb592-6bd8-4b57-860d-52b34c6cbbb1"
AT_TYPE_KEY = "location_type_AT_Automotive_Testing"
NUERNBERG = (49.4521, 11.0767)
NUERNBERG_KM = 30
QUOTA = {"S1": 12, "S2": 12, "S3": 8, "S4": 10, "S5": 7, "S6": 7, "S7": 6, "S8": 5, "S9": 6}

S1_NAME_WORDS = ["station", "standort", "niederlassung", "außenstelle", "prüfstelle", "gutachten", "ga-büro", "automobil"]
SITE_PREFIX = re.compile(r"(station|standort|niederlassung|außenstelle|(kfz-)?prüfstelle|ga-büro|gutachten\w*)\b", re.I)
PERSON_TOKEN = re.compile(r"\bDipl|\bIng\b|\bInh\.|Sachverständig|Ingenieur|\bB\.Eng\.|\bM\.Sc\.|\bDr\.|\bund\b", re.I)
CAP_WORD = re.compile(r"^[A-ZÄÖÜ][\w\-\.]*$")
COMPANY_SUFFIX = {"gmbh", "ag", "kg", "gbr", "ug", "mbh", "ohg", "e.k.", "co.", "se", "e.v."}
HU_WORD = re.compile(r"\bHU\b")
HU_LONG = re.compile(r"hauptuntersuchung", re.I)

COLUMNS = [
    "sample_id", "stratum", "id", "name", "name2", "street", "zip", "city", "state", "types", "nfz",
    "booking_link", "hu_mention", "special_services", "website_url", "maps_url", "freetext_excerpt",
    "same_address_count",
    "hu_on_site", "site_kind", "evidence", "labelled_by", "labelled_on",
]
LABEL_COLUMNS = COLUMNS[-5:]
URL_COLUMNS = ("booking_link", "website_url", "maps_url")
TEXT_COLUMNS = ("sample_id", "id", "zip")


# ---------------------------------------------------------------- record helpers

def type_ids(r):
    return [t["id"] for t in r["locationTypesResolved"]]


def type_names(r):
    return [t["name"] for t in r["locationTypesResolved"]]


def norm_zip(r) -> str:
    z = (r.get("zipCode") or "").strip()
    return z.zfill(5) if r["countryResolved"].get("code") == "DE" else z.replace(" ", "")


def hu_fields(r) -> list[str]:
    hits = []
    for f in recon.TEXT_FIELDS + ["name", "name2"]:
        txt = " ".join(recon.strings(r.get(f)))
        if HU_LONG.search(txt) or HU_WORD.search(txt):
            hits.append(f)
    return hits


def booking_link(r) -> str:
    link = r.get("linkOnlineAppointment")
    return (link.get("to") or "").strip() if not recon.is_empty(link) else ""


def rich_text(v) -> str:
    """Flatten a CMS RichText value to plain text (text leaves only, paragraphs joined by a space)."""
    out: list[str] = []

    def rec(x):
        if isinstance(x, str):
            out.append(x)
        elif isinstance(x, dict):
            for key in ("elements", "content"):
                if key in x:
                    rec(x[key])
        elif isinstance(x, list):
            for y in x:
                rec(y)

    rec(v)
    return re.sub(r"\s+", " ", " ".join(out)).strip()


def street_key(r) -> str:
    s = f"{r.get('street') or ''} {r.get('houseNumber') or ''}".lower().replace("ß", "ss")
    s = re.sub(r"stra(ss|ß)e|str\.", "str", s)
    return re.sub(r"[^a-z0-9äöü]", "", s)


def website_url(r) -> tuple[str, str]:
    """(url, kind) with kind in dekra.de / other_domain / search_fallback."""
    link = r.get("linkWebsite")
    to = ((link or {}).get("to") or "").strip() if isinstance(link, dict) else ""
    if to.startswith("/"):
        to = "https://www.dekra.de" + to
    if to.startswith("http"):
        kind = "dekra.de" if re.match(r"https?://(www\.)?dekra\.de/", to) else "other_domain"
        return to, kind
    q = quote_plus(f"DEKRA {r.get('name', '')} {r.get('city', '')}".replace("\n", " "))
    return f"https://www.google.com/search?q={q}", "search_fallback"


# ---------------------------------------------------------------- strata

def person_rule(name: str) -> str | None:
    """S2 heuristic. Returns the rule that matched ('token' or 'caps') or None."""
    n = name.strip()
    if "dekra" in n.lower() or SITE_PREFIX.match(n):
        return None
    if PERSON_TOKEN.search(n):
        return "token"
    words = n.split()
    if 2 <= len(words) <= 3 and all(CAP_WORD.match(w) for w in words) \
            and not any(w.lower() in COMPANY_SUFFIX for w in words) and not re.search(r"\d", n):
        return "caps"
    return None


def auto_stratum(r) -> str:
    """Natural stratum S2..S9 of an Automobil-typed record, first match in order."""
    n = r["name"].strip()
    if person_rule(n):
        return "S2"
    low = n.lower()
    if re.match(r"station\b", low):
        b, h = bool(booking_link(r)), bool(hu_fields(r))
        return "S3" if b and h else "S4" if not b and not h else "S9"
    if re.match(r"(kfz-)?prüfstelle\b", low):
        return "S5"
    if re.match(r"niederlassung\b", low):
        return "S6"
    if re.match(r"außenstelle\b", low):
        return "S7"
    if re.match(r"standort\b", low):
        return "S8"
    return "S9"


def pick(cands: list[dict], quota: int | None, state_of, must: list[dict] | None = None) -> list[dict]:
    """Take `must` first, then fill up to `quota` round-robin across Bundesländer (seed 42). quota None = all."""
    chosen = list(must or [])
    if quota is None:
        return chosen + [r for r in sorted(cands, key=lambda r: r["id"]) if r not in chosen]
    rng = random.Random(SEED)
    rest = [r for r in sorted(cands, key=lambda r: r["id"]) if r not in chosen]
    groups: dict[str, list[dict]] = defaultdict(list)
    for r in rest:
        groups[state_of(r)].append(r)
    order = sorted(groups)
    rng.shuffle(order)
    for g in groups.values():
        rng.shuffle(g)
    while len(chosen) < quota and any(groups[s] for s in order):
        for s in order:
            if groups[s] and len(chosen) < quota:
                chosen.append(groups[s].pop())
    return chosen


# ---------------------------------------------------------------- main

def main(run_name: str = DEFAULT_RUN) -> None:
    run = paths.scraper_dir("dekra") / "test_output" / run_name
    universe = recon.load_responses(run)["R09"]
    print(f"R09: {len(universe)} records from {run.name}")

    plz = load_plz_centroids(paths.geonames_dir() / "DE.txt").set_index("zip")["state"].to_dict()

    def state_of(r) -> str:
        return plz.get(norm_zip(r), "") or "(unbekannt)"

    addr_count = Counter((street_key(r), norm_zip(r)) for r in universe)

    de = [r for r in universe if r["countryResolved"].get("code") == "DE"]
    pure_whg = [r for r in de if type_names(r) == ["WHG Standort"]]
    pop = [r for r in de if type_names(r) != ["WHG Standort"]]
    print(f"DE {len(de)}, pure WHG {len(pure_whg)}, population {len(pop)}")

    # ---- S1
    def is_s1(r) -> bool:
        untyped_or_at = not r["locationTypesResolved"] or any(t["key"] == AT_TYPE_KEY for t in r["locationTypesResolved"])
        return untyped_or_at and any(w in r["name"].lower() for w in S1_NAME_WORDS)

    s1_all = [r for r in pop if is_s1(r)]
    auto = [r for r in pop if AUTO_TYPE_ID in type_ids(r)]

    # ---- S10 (set aside first so nothing is sampled twice)
    def near_nuernberg(r) -> bool:
        return float(haversine_km(NUERNBERG[0], NUERNBERG[1], r["latitude"], r["longitude"])) <= NUERNBERG_KM

    s10 = sorted((r for r in auto if near_nuernberg(r)), key=lambda r: (state_of(r), r["name"], r["id"]))
    s10_ids = {r["id"] for r in s10}

    natural = {r["id"]: auto_stratum(r) for r in auto}
    pools: dict[str, list[dict]] = defaultdict(list)
    for r in auto:
        if r["id"] not in s10_ids:
            pools[natural[r["id"]]].append(r)

    sample: list[tuple[str, dict]] = []
    sample += [("S1", r) for r in pick(s1_all, QUOTA["S1"], state_of)]
    for s in ("S2", "S3", "S4"):
        sample += [(s, r) for r in pick(pools[s], QUOTA[s], state_of)]
    s5_must = [r for r in sorted(pools["S5"], key=lambda r: r["id"]) if not booking_link(r) and not hu_fields(r)]
    sample += [("S5", r) for r in pick(pools["S5"], QUOTA["S5"], state_of, must=s5_must)]
    for s in ("S6", "S7", "S8", "S9"):
        sample += [(s, r) for r in pick(pools[s], QUOTA[s], state_of)]
    sample += [("S10", r) for r in s10]
    assert len({r["id"] for _, r in sample}) == len(sample), "a record was sampled twice"

    # ---- rows
    rows = []
    site_kinds = Counter()
    for n, (s, r) in enumerate(sample, start=1):
        url, kind = website_url(r)
        site_kinds[kind] += 1
        street = f"{r.get('street') or ''} {r.get('houseNumber') or ''}".strip()
        rows.append({
            "sample_id": f"DEKRA-{n:03d}", "stratum": s, "id": r["id"],
            "name": r["name"].strip(), "name2": (r.get("name2") or "").strip(),
            "street": re.sub(r"\s+", " ", street), "zip": norm_zip(r), "city": (r.get("city") or "").strip(),
            "state": state_of(r), "types": " | ".join(type_names(r)) or "(none)",
            "nfz": "yes" if r["suitableForCommercialVehicles"] else "no",
            "booking_link": booking_link(r), "hu_mention": "; ".join(hu_fields(r)),
            "special_services": "; ".join(o["value"] for o in r["specialServicesList"]),
            "website_url": url,
            "maps_url": f"https://www.google.com/maps?q={r['latitude']},{r['longitude']}",
            "freetext_excerpt": rich_text(r.get("freeText1"))[:300],
            "same_address_count": addr_count[(street_key(r), norm_zip(r))],
            **{c: "" for c in LABEL_COLUMNS},
        })
    sample_id_of = {row["id"]: row["sample_id"] for row in rows}

    # ---- S2 heuristic match list (all Automobil-typed DE records, including those set aside for S10)
    s2_matches = [(person_rule(r["name"]), r) for r in auto if person_rule(r["name"])]

    # ---- outputs
    today = date.today().isoformat()
    out_dir = paths.reference_dir() / "ground_truth"
    out_dir.mkdir(parents=True, exist_ok=True)
    suffix = ""
    k = 1
    while any((out_dir / f"{stem}_{today}{suffix}{ext}").exists()
              for stem, ext in (("dekra_labelling_sample", ".xlsx"), ("dekra_labelling_sample", ".csv"),
                                ("dekra_labelling_s2_matches", ".csv"))):
        k += 1
        suffix = f"_{k}"
    xlsx = out_dir / f"dekra_labelling_sample_{today}{suffix}.xlsx"
    csv_path = out_dir / f"dekra_labelling_sample_{today}{suffix}.csv"
    s2_path = out_dir / f"dekra_labelling_s2_matches_{today}{suffix}.csv"

    with open(csv_path, "x", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
    with open(s2_path, "x", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["rule", "sampled_as", "id", "name", "name2", "zip", "city", "state", "set_aside_for_s10"])
        for rule, r in sorted(s2_matches, key=lambda x: (x[0], x[1]["name"])):
            w.writerow([rule, sample_id_of.get(r["id"], ""), r["id"], r["name"].strip(), (r.get("name2") or "").strip(),
                        norm_zip(r), r.get("city"), state_of(r), "yes" if r["id"] in s10_ids else ""])
    write_xlsx(xlsx, rows)

    # ---- console summary (counts only; no personal names)
    print("\nstratum sizes (population N / sampled n)")
    for s in ["S1"] + [f"S{i}" for i in range(2, 10)] + ["S10"]:
        big_n = (len(s1_all) if s == "S1" else len(s10) if s == "S10" else len(pools[s]))
        print(f"  {s:4} N={big_n:4} n={sum(1 for x, _ in sample if x == s)}")
    print(f"S5 'neither signal' forced in: {len(s5_must)}")
    print(f"natural stratum of the S10 records: {dict(Counter(natural[i] for i in s10_ids))}")
    print(f"S2 matches: {Counter(rule for rule, _ in s2_matches)}; of which set aside for S10: "
          f"{sum(1 for _, r in s2_matches if r['id'] in s10_ids)}")
    print(f"website_url kinds in sample: {dict(site_kinds)}")
    print(f"states in sample: {dict(Counter(row['state'] for row in rows).most_common())}")
    print(f"S1 candidate states: {dict(Counter(state_of(r) for r in s1_all).most_common())}")
    print(f"same_address_count>1 in sample: {sum(1 for row in rows if row['same_address_count'] > 1)}")
    print(f"total rows {len(rows)}")
    for p in (xlsx, csv_path, s2_path):
        print("wrote", p)


GUIDE = [
    ("DEKRA labelling sample: guide", None),
    ("DRAFT. The labelling rules that should follow the request were cut off and did not arrive. "
     "The text below only restates the column meanings and allowed values; replace it with the agreed rules.", None),
    ("", None),
    ("What to decide for each row", None),
    ("hu_on_site", "yes / no / unclear: can a driver have a Hauptuntersuchung (HU) done at this address?"),
    ("site_kind", "pruefstelle: DEKRA inspection site open to the public. "
                  "gutachten_office: expert-report / damage-assessment office without HU testing. "
                  "partner_engineer: independent engineer or engineering office working with DEKRA (often person-named). "
                  "admin_or_other_division: head office, academy, labour, certification or another DEKRA division. "
                  "other: anything else."),
    ("evidence", "Short free text: what you saw and where (website page, booking page, Google Maps, phone)."),
    ("labelled_by / labelled_on", "Your name and the date (YYYY-MM-DD)."),
    ("", None),
    ("Where the columns come from (DEKRA API, recon run 2026-09-30_1831_recon)", None),
    ("stratum", "S1 untyped or AT-type 'Automobil Prüfwesen' with a vehicle-site name; S2 person-named; S3 Station with booking link and HU mention; "
                "S4 Station with neither; S5 Kfz-Prüfstelle / Prüfstelle; S6 Niederlassung; S7 Außenstelle; S8 Standort; S9 everything else "
                "(incl. Stations with only one signal); S10 every Automobil-typed record within 30 km of Nürnberg."),
    ("booking_link", "linkOnlineAppointment.to, full URL, empty if none."),
    ("hu_mention", "Fields whose text contains 'Hauptuntersuchung' or the word 'HU' (freeText1, freeText2, services, currentInformation, openingHoursAreas, name, name2)."),
    ("website_url", "linkWebsite.to. If a record has no website link: a Google search URL 'DEKRA <name> <city>'."),
    ("same_address_count", "Records in the whole universe (all types, all countries) with the same normalised street + house number + PLZ, including this one. 1 = only this record."),
    ("nfz", "suitableForCommercialVehicles."),
    ("state", "Bundesland from the GeoNames PLZ file (CC BY 4.0), not from the DEKRA record."),
]


def write_xlsx(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sample"
    ws.append(COLUMNS)
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    label_fill = PatternFill("solid", fgColor="FFF2CC")
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = label_fill if name in LABEL_COLUMNS else head_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            value = row[name]
            cell.value = value if value != "" else None
            if name in TEXT_COLUMNS:
                cell.value = str(value)
                cell.number_format = "@"
                cell.data_type = "s"
            if name in URL_COLUMNS and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
            if name in ("freetext_excerpt", "evidence", "special_services", "types", "hu_mention"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
    widths = {"sample_id": 11, "stratum": 8, "id": 38, "name": 34, "name2": 22, "street": 26, "zip": 7, "city": 18, "state": 20,
              "types": 28, "nfz": 5, "booking_link": 30, "hu_mention": 20, "special_services": 28, "website_url": 30, "maps_url": 26,
              "freetext_excerpt": 50, "same_address_count": 9, "hu_on_site": 11, "site_kind": 22, "evidence": 40, "labelled_by": 12,
              "labelled_on": 12}
    for c_idx, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=c_idx).column_letter].width = widths[name]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    last = len(rows) + 1
    dv_hu = DataValidation(type="list", formula1='"yes,no,unclear"', allow_blank=True)
    dv_kind = DataValidation(type="list", formula1='"pruefstelle,gutachten_office,partner_engineer,admin_or_other_division,other"', allow_blank=True)
    ws.add_data_validation(dv_hu)
    ws.add_data_validation(dv_kind)
    dv_hu.add(f"{ws.cell(row=2, column=COLUMNS.index('hu_on_site') + 1).column_letter}2:"
              f"{ws.cell(row=2, column=COLUMNS.index('hu_on_site') + 1).column_letter}{last}")
    dv_kind.add(f"{ws.cell(row=2, column=COLUMNS.index('site_kind') + 1).column_letter}2:"
                f"{ws.cell(row=2, column=COLUMNS.index('site_kind') + 1).column_letter}{last}")

    guide = wb.create_sheet("Guide")
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 120
    for a, b in GUIDE:
        if b is None:
            guide.append([a])
            guide.cell(row=guide.max_row, column=1).font = Font(bold=True)
        else:
            guide.append([a, b])
            guide.cell(row=guide.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            guide.cell(row=guide.max_row, column=1).alignment = Alignment(vertical="top")
    wb.save(path)  # path is new by construction (checked above), so this never overwrites


if __name__ == "__main__":
    main(*(sys.argv[1:2]))
