"""Parser tests against tests/fixtures/response_muenchen_50km_2026-09-28.json (München 80686, radius 50 km).

The fixture is anonymised (2026-10-28): names, phones, emails and URLs are fake ("Muster Prüfbüro NN"); ids,
addresses, coordinates and branches are the real ones. The original is in 4_reference/api_samples/gtue/ (not in git).
"""

import pytest

from gtue.adapter import classify_site_type, export, load_response, parse_response, validate

from conftest import SAMPLE


@pytest.fixture(scope="module")
def df():
    if not SAMPLE.exists():
        pytest.skip(f"Sample file missing: {SAMPLE}")
    return parse_response(load_response(SAMPLE), scraped_at="2026-09-28T11:26:10")


def test_row_count_and_unique_ids(df):
    assert len(df) == 20
    assert df["source_id"].is_unique


def test_coordinates_in_germany(df):
    # Catches X/Y swaps: Germany is lat ~47-55, lng ~6-15
    assert df["lat"].between(47, 56).all()
    assert df["lng"].between(5, 16).all()


def test_site_type_counts(df):
    counts = df["site_type"].value_counts().to_dict()
    assert counts.get("pruefstelle", 0) == 13
    assert counts.get("partner_office", 0) == 7
    assert counts.get("unknown", 0) == 0
    # scope rule 2026-10-07: the 7 offices count when they list HU (6 of them do), see test_gtue_scope.py
    assert df["in_scope"].sum() == 19


def test_location_id_format(df):
    assert (df["location_id"] == "GTUE-" + df["source_id"]).all()


def test_source_id_keeps_leading_zeros(df):
    assert "001241101" in df["source_id"].tolist()


def test_trudering_record(df):
    row = df.loc[df["source_id"] == "166591000"].iloc[0]
    assert row["name"] == "Muster Prüfbüro 07"
    assert row["zip"] == "81827"
    assert row["city"] == "München"
    assert row["site_type"] == "pruefstelle"


def test_name_addon2_aubing(df):
    row = df.loc[df["source_id"] == "176883000"].iloc[0]
    assert row["name"] == "Muster Prüfbüro 02"
    assert row["name_addon2"] == "Prüfstelle Aubing"


def test_gtup_code_matches_text(df):
    # Code GTUP is the primary signal; on the sample it agrees with the text on every row
    assert df["inspects_on_site"].sum() == 13
    assert df["site_type_evidence"].str.contains("code:GTUP; text:yes").sum() == 13
    assert df["site_type_evidence"].str.contains("code:none; text:no").sum() == 7
    row = df.loc[df["source_id"] == "166591000"].iloc[0]
    assert row["site_type_evidence"] == "herkunft:E; code:GTUP; text:yes"


@pytest.mark.parametrize(
    "herkunft, has_code, has_text, expected",
    [
        ("E", True, True, "pruefstelle"),
        ("G", False, False, "partner_office"),
        ("G", True, True, "unknown"),  # Herkunft disagrees with vor Ort
        ("E", False, False, "unknown"),
        ("E", True, False, "unknown"),  # code and text disagree
        ("E", False, True, "unknown"),
        (None, True, True, "unknown"),  # Herkunft missing
    ],
)
def test_classify_site_type(herkunft, has_code, has_text, expected):
    site_type, _, needs_review = classify_site_type(herkunft, has_code, has_text)
    assert site_type == expected
    assert needs_review == (expected == "unknown")


def test_id_prefix_check(df):
    report = validate(df, load_response(SAMPLE))
    assert report["id_prefix_match_count"] == 20
    assert report["id_prefix_breaks"] == []


def test_xlsx_export(df, tmp_path):
    from openpyxl import load_workbook

    written = export(df, tmp_path, run_date="2026-09-28")
    path, rows = written["pruefstellen_xlsx"]
    assert rows == 19
    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    assert ws.freeze_panes == "A2"

    all_ws = load_workbook(written["all_xlsx"][0]).active
    ids = [c.value for c in all_ws["C"][1:]]  # source_id column
    assert header.index("source_id") == 2
    assert "001241101" in ids  # leading zeros kept
    for col in ("source_id", "location_id", "zip"):
        cell = all_ws.cell(row=2, column=header.index(col) + 1)
        assert isinstance(cell.value, str) and cell.number_format == "@"
    for col in ("lat", "lng"):
        assert isinstance(all_ws.cell(row=2, column=header.index(col) + 1).value, float)
    assert all_ws.column_dimensions["A"].width > 10
