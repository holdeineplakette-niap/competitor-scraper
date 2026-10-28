"""GTÜ source adapter.

Parses responses of the GTÜ finder API (SmartMaps, endpoint
`addresssearch/ByBranchesAndAddress`), classifies each location's site type,
validates a parsed batch, exports CSV and XLSX files, and fetches the whole
finder with the adaptive grid in core/grid.py.

Parsing, classification and validation are pure functions. Network access
happens only in GtueClient; run(mode="simulated") uses an offline fake API.

Every run writes one folder (see core/paths.py and
0_instructions/storage_and_naming_convention.md): raw.jsonl, query_log.csv,
run_summary.json/.md, the exports and, if squares stayed open, open_squares.csv.
"""

from __future__ import annotations

import json
import logging
import math
import os
import random
import re
import time
from dataclasses import dataclass
from datetime import date, datetime
from itertools import islice
from pathlib import Path

import pandas as pd
import requests
from dotenv import dotenv_values
from openpyxl.utils import get_column_letter

from core import enrich, paths
from core.geo import ensure_geonames, haversine_km, load_plz_centroids
from core.grid import (
    BIG_CITIES,
    LOG_COLUMNS,
    GridResult,
    QueryPoint,
    QueryResult,
    SimulatedFinder,
    StopRun,
    fake_locations,
    geocoding_outliers,
    cap_probe,
    reference_check,
    reference_recheck,
    run_grid,
)

logger = logging.getLogger(__name__)

SOURCE_ORG = "GTÜ"
ID_PREFIX = "GTUE-"

# BranchListElements entry "GTÜ-Ingenieurbüro mit amtlicher Fahrzeugprüfung vor Ort":
# the code is the primary signal, the text a cross-check.
VOR_ORT_CODE = "GTUP"
VOR_ORT_TEXT = "Fahrzeugprüfung vor Ort"
HU_HEADLINE = "Hauptuntersuchung Pkw"
REASON_SITE_TYPE_UNKNOWN = "site_type unknown"

# Scope rule (Nikos, 2026-10-07): "if we do not have certain cause for exclusion, we include them".
# The finder's three signals (herkunft, code, text) say whether the office has an own on-site test
# station, but hand checks showed that G offices which list HU do HU (5 of 5), while G offices
# without any HU service do not (5 of 5, e.g. a lift inspector). So the service list decides.
TIER_PRUEFSTELLE = "pruefstelle"  # herkunft E + code GTUP + text: own on-site inspection
TIER_UNKNOWN = "unknown_conflict"  # the three signals disagree: in scope, needs_review
TIER_OFFICE_HU = "office_hu_listed"  # G office (no on-site code) that lists Hauptuntersuchung
TIER_OFFICE_NO_HU = "office_no_hu"  # G office without any HU service: out of scope
SCOPE_TIERS = (TIER_PRUEFSTELLE, TIER_UNKNOWN, TIER_OFFICE_HU, TIER_OFFICE_NO_HU)
IN_SCOPE_TIERS = (TIER_PRUEFSTELLE, TIER_UNKNOWN, TIER_OFFICE_HU)
REASON_NO_HU = "no HU offered (finder service list)"

# Rows at one physical site: same normalised street + PLZ, or coordinates this close (km).
SAME_SITE_KM = 0.15

# The API returns locations abroad too (IsoCountryCode doesn't filter). They stay
# in the all-files but are never in scope.
HOME_COUNTRY = "D"  # Address.Country value for Germany
REASON_OUTSIDE_GERMANY = "outside Germany"
REASON_COUNTRY_MISSING = "country missing"

_BR = re.compile(r"\s*<br\s*/?>\s*", re.IGNORECASE)

COLUMNS = [
    "location_id",
    "source_org",
    "source_id",
    "name",
    "name_addon",
    "name_addon2",
    "street",
    "zip",
    "city",
    "state",
    "country",
    "lat",
    "lng",
    "phone",
    "mobile",
    "email",
    "website",
    "gtue_profile_url",
    "booking_url",
    "opening_hours",
    "services",
    "offers_hu",
    "inspects_on_site",
    "herkunft",
    "site_type",
    "site_type_evidence",
    "needs_review",
    "needs_review_reason",
    "in_scope",
    "scope_tier",
    "scope_reason",
    "source_data_date",
    "source_last_change",
    "scraped_at",
]

# Added after parsing, by enrich_german (needs all rows at once)
SAME_SITE_COLUMN = "same_site_ids"


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def _clean(value) -> str | None:
    """Strip strings; empty strings and None become None."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _to_float(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _memos(basic: dict, memo_type: str) -> list[dict]:
    return [m for m in basic.get("MemoItems") or [] if m.get("Type") == memo_type]


def _parse_json_info(basic: dict) -> tuple[str | None, str | None]:
    """Return (Herkunft, Stand) from the JSON_INFO memo. Missing or invalid JSON -> None."""
    for memo in _memos(basic, "JSON_INFO"):
        try:
            info = json.loads(memo.get("Text") or "")
        except (TypeError, ValueError):
            continue
        hidden = info.get("HiddenPoiInfos") if isinstance(info, dict) else None
        if isinstance(hidden, dict):
            return _clean(hidden.get("Herkunft")), _clean(hidden.get("Stand"))
    return None, None


def _opening_hours(basic: dict) -> str | None:
    parts = [_clean(_BR.sub(" | ", m.get("Text") or "")) for m in _memos(basic, "OPENING_HOURS")]
    parts = [p for p in parts if p]
    return " | ".join(parts) or None


def _booking_url(basic: dict) -> str | None:
    for memo in _memos(basic, "CUSTOM"):
        if memo.get("Title") == "TerminUrl" and _clean(memo.get("Text")):
            return _clean(memo.get("Text"))
    return None


# --------------------------------------------------------------------------
# Site type
# --------------------------------------------------------------------------

def classify_site_type(herkunft: str | None, has_code: bool, has_text: bool) -> tuple[str, str, bool]:
    """Return (site_type, site_type_evidence, needs_review).

    has_code: branch code GTUP present; has_text: "Fahrzeugprüfung vor Ort" text present.
    All three signals must agree; any disagreement or missing value -> "unknown".
    """
    evidence = (
        f"herkunft:{herkunft or 'missing'}; "
        f"code:{VOR_ORT_CODE if has_code else 'none'}; "
        f"text:{'yes' if has_text else 'no'}"
    )
    if has_code != has_text:
        return "unknown", evidence, True
    if has_code and herkunft == "E":
        return "pruefstelle", evidence, False
    if not has_code and herkunft == "G":
        return "partner_office", evidence, False
    return "unknown", evidence, True


def classify_scope(site_type: str, offers_hu: bool) -> tuple[str, bool, str | None]:
    """Return (scope_tier, in_scope before the country check, reason if out of scope).

    Only a certain cause excludes a row: a G office whose service list has no HU at all.
    """
    if site_type == "pruefstelle":
        return TIER_PRUEFSTELLE, True, None
    if site_type == "partner_office":
        if offers_hu:
            return TIER_OFFICE_HU, True, None
        return TIER_OFFICE_NO_HU, False, REASON_NO_HU
    return TIER_UNKNOWN, True, None  # signals disagree: keep, needs_review is already set


# --------------------------------------------------------------------------
# Parser
# --------------------------------------------------------------------------

def parse_item(item: dict, scraped_at: str) -> dict:
    """Parse one entry of AddressItems into a flat row."""
    basic = item.get("BasicData") or {}
    ids = basic.get("Identifiers") or {}
    address = basic.get("Address") or {}
    geo = basic.get("Geo") or {}
    contact = basic.get("Contact") or {}
    business = basic.get("BusinessData") or {}

    source_id = _clean(ids.get("ProviderForeignKey"))
    herkunft, stand = _parse_json_info(basic)

    branches = basic.get("BranchListElements") or []
    has_code = any(b.get("BranchCode") == VOR_ORT_CODE for b in branches)
    has_text = any(VOR_ORT_TEXT in (b.get("BranchText") or "") for b in branches)

    headlines = [h for h in (_clean(o.get("Headline")) for o in basic.get("ObjectListItems") or []) if h]

    site_type, evidence, needs_review = classify_site_type(herkunft, has_code, has_text)
    offers_hu = HU_HEADLINE in headlines
    scope_tier, tier_in_scope, scope_reason = classify_scope(site_type, offers_hu)

    country = _clean(address.get("Country"))
    in_germany = country == HOME_COUNTRY
    reasons = [REASON_SITE_TYPE_UNKNOWN] if needs_review else []
    if not in_germany:
        scope_reason = REASON_OUTSIDE_GERMANY if country else REASON_COUNTRY_MISSING
        # Explains why the row is out of scope. It goes here, not in site_type_evidence
        # (which only records the classification signals, unchanged abroad); needs_review
        # itself is not set, because the exclusion is certain, not a doubt.
        reasons.append(REASON_OUTSIDE_GERMANY if country else REASON_COUNTRY_MISSING)

    return {
        "location_id": f"{ID_PREFIX}{source_id}" if source_id else None,
        "source_org": SOURCE_ORG,
        "source_id": source_id,
        "name": _clean(address.get("CompanyName")),
        "name_addon": _clean(address.get("CompanyNameAddon1")),
        "name_addon2": _clean(address.get("CompanyNameAddon2")),  # site label, e.g. "Prüfstelle Aubing"
        "street": _clean(address.get("Street")),
        "zip": _clean(address.get("Zip")),
        "city": _clean(address.get("City")),
        "state": None,  # not in the record; filled from the PLZ by core.enrich
        "country": country,
        "lat": _to_float(geo.get("YCoord")),  # Y = latitude
        "lng": _to_float(geo.get("XCoord")),  # X = longitude
        "phone": _clean(contact.get("Phone")),
        "mobile": _clean(contact.get("Mobile")),
        "email": _clean(contact.get("Email")),
        "website": _clean(contact.get("Url")),
        "gtue_profile_url": _clean(contact.get("Url3")),
        "booking_url": _booking_url(basic),
        "opening_hours": _opening_hours(basic),
        "services": "; ".join(headlines) or None,
        "offers_hu": offers_hu,
        "inspects_on_site": has_code,
        "herkunft": herkunft,
        "site_type": site_type,
        "site_type_evidence": evidence,
        "needs_review": needs_review,
        "needs_review_reason": enrich.REASON_SEPARATOR.join(reasons) or None,
        "in_scope": tier_in_scope and in_germany,
        "scope_tier": scope_tier,
        "scope_reason": scope_reason,
        "source_data_date": stand,
        # DateLastChange is a batch timestamp (identical on every record, e.g.
        # "2026-09-21 08:30"), not a per-location change date. Do NOT use it for
        # change detection; diff our own snapshots instead.
        "source_last_change": _clean(business.get("DateLastChange")),
        "scraped_at": scraped_at,
    }


def parse_response(data: dict, scraped_at: str) -> pd.DataFrame:
    """One row per entry in data["AddressItems"]. Geo.Distance is deliberately not stored."""
    rows = [parse_item(item, scraped_at) for item in data.get("AddressItems") or []]
    return pd.DataFrame(rows, columns=COLUMNS)


def load_response(path: str | Path) -> dict:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _address_key(street, zip_) -> str | None:
    """Normalised street + PLZ, or None if either is missing."""
    if not street or not zip_:
        return None
    s = str(street).lower().replace("ß", "ss")
    s = re.sub(r"stra(ss|ß)e|str\.?", "str", s)
    s = re.sub(r"[^a-z0-9]", "", s)
    return f"{str(zip_).strip()}|{s}" if s else None


def add_same_site_ids(df: pd.DataFrame) -> pd.DataFrame:
    """Add `same_site_ids`: the other German rows at the same site (same normalised street + PLZ, or within
    SAME_SITE_KM by coordinates), joined with the reason separator. Information only: nothing is merged or
    dropped here. GTÜ lists one site twice (an E and a G listing) in about 100 cases; merging is the
    aggregator's job (core.dedupe, strong evidence), so this column just says where to look.
    """
    out = df.copy()
    out[SAME_SITE_COLUMN] = None
    german = out.index[out["country"].eq(HOME_COUNTRY)].tolist()
    if len(german) < 2:
        return out
    ids = out.loc[german, "location_id"].tolist()
    keys = [_address_key(s, z) for s, z in zip(out.loc[german, "street"], out.loc[german, "zip"])]
    lat = out.loc[german, "lat"].astype(float).to_numpy()
    lng = out.loc[german, "lng"].astype(float).to_numpy()
    by_key: dict[str, list[int]] = {}
    for i, k in enumerate(keys):
        if k:
            by_key.setdefault(k, []).append(i)
    found: list[list[str]] = []
    for i in range(len(german)):
        same = {j for j in by_key.get(keys[i], []) if j != i} if keys[i] else set()
        if lat[i] == lat[i] and lng[i] == lng[i]:  # not NaN
            dy = (lat - lat[i]) * 111.0
            dx = (lng - lng[i]) * 111.0 * math.cos(math.radians(lat[i]))
            near = (dy * dy + dx * dx) <= SAME_SITE_KM**2  # NaN compares False
            same |= {int(j) for j in near.nonzero()[0] if j != i}
        found.append(sorted(ids[j] for j in same if ids[j]))
    out.loc[german, SAME_SITE_COLUMN] = [enrich.REASON_SEPARATOR.join(f) or None for f in found]
    return out


def enrich_german(df: pd.DataFrame, centroids: pd.DataFrame) -> pd.DataFrame:
    """core.enrich for German rows only (state, PLZ distance, PLZ flags), then the same-site information."""
    return add_same_site_ids(enrich.enrich(df, centroids, check=df["country"].eq(HOME_COUNTRY)))


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------

def _has_reason(df: pd.DataFrame, reason: str) -> pd.Series:
    if "needs_review_reason" not in df.columns:
        return pd.Series(False, index=df.index)
    return df["needs_review_reason"].map(lambda v: reason in enrich.split_reasons(v)).astype(bool)


def validate(df: pd.DataFrame, data: dict) -> dict:
    """Check a parsed (and ideally enriched) response for completeness and data quality."""
    paging = data.get("Paging") or {}
    count = paging.get("Count")
    max_count = paging.get("MaxCount")
    warnings = []

    if count is not None and len(df) != count:
        warnings.append(f"Rows parsed ({len(df)}) != Paging.Count ({count})")
    if count is not None and max_count is not None and count < max_count:
        warnings.append(
            f"Paging.Count ({count}) < MaxCount ({max_count}): the query did not return everything"
        )

    dup_mask = df["source_id"].duplicated(keep=False)
    duplicate_ids = sorted(df.loc[dup_mask, "source_id"].dropna().unique().tolist())
    if duplicate_ids:
        warnings.append(f"Duplicate source_ids: {duplicate_ids}")

    missing_coords = df.loc[df["lat"].isna() | df["lng"].isna(), "source_id"].tolist()
    if missing_coords:
        warnings.append(f"Rows missing lat/lng: {missing_coords}")
    missing_zip = df.loc[df["zip"].isna(), "source_id"].tolist()
    if missing_zip:
        warnings.append(f"Rows missing zip: {missing_zip}")

    needs_review = int(df["needs_review"].sum())
    if needs_review:
        warnings.append(f"{needs_review} rows need review")

    # PLZ checks from core.enrich (only if enrichment was applied)
    plz_checked = "plz_distance_km" in df.columns
    far = df.loc[_has_reason(df, enrich.REASON_FAR_FROM_PLZ), ["source_id", "zip", "plz_distance_km"]] if plz_checked else df.iloc[0:0]
    unknown = df.loc[_has_reason(df, enrich.REASON_UNKNOWN_PLZ), ["source_id", "zip"]] if plz_checked else df.iloc[0:0]

    # Observed pattern (report only, not used for classification):
    # Herkunft E -> source_id starts with "1", Herkunft G -> starts with "00"
    sid = df["source_id"].fillna("")
    fits = ((df["herkunft"] == "E") & sid.str.startswith("1")) | (
        (df["herkunft"] == "G") & sid.str.startswith("00")
    )
    fits = fits.fillna(False).astype(bool)
    id_prefix_breaks = df.loc[~fits, ["source_id", "herkunft", "name"]].to_dict("records")

    return {
        "rows_parsed": len(df),
        "paging_count": count,
        "paging_max_count": max_count,
        "complete": count is not None and count == max_count,
        "site_type_counts": df["site_type"].value_counts().to_dict(),
        "in_scope_count": int(df["in_scope"].sum()),
        "needs_review_count": needs_review,
        "needs_review_by_reason": enrich.reason_counts(df),
        "fill_rate_pct": {c: round(df[c].notna().mean() * 100, 1) if len(df) else 0.0 for c in df.columns},
        "duplicate_source_ids": duplicate_ids,
        "rows_missing_lat_lng": missing_coords,
        "rows_missing_zip": missing_zip,
        "plz_check_applied": plz_checked,
        "max_plz_distance_km": float(df["plz_distance_km"].max()) if plz_checked and df["plz_distance_km"].notna().any() else None,
        "rows_far_from_plz": far.to_dict("records"),
        "rows_unknown_plz": unknown.to_dict("records"),
        "id_prefix_match_count": int(fits.sum()),
        "id_prefix_breaks": id_prefix_breaks,
        "warnings": warnings,
    }


def format_report(report: dict) -> str:
    """Readable text version of a validate() report."""
    lines = [
        "GTÜ validation report",
        "=" * 40,
        f"Rows parsed:        {report['rows_parsed']}",
        f"Paging.Count:       {report['paging_count']}",
        f"Paging.MaxCount:    {report['paging_max_count']}",
        f"Complete query:     {report['complete']}",
        "",
        "Site types:",
    ]
    lines += [f"  {k:<16}{v}" for k, v in report["site_type_counts"].items()]
    lines += [
        f"In scope:           {report['in_scope_count']}",
        f"Needs review:       {report['needs_review_count']}",
        "Reasons (needs_review_reason):" if report["needs_review_by_reason"] else "Reasons:           none",
    ]
    lines += [f"  {k:<26}{v}" for k, v in report["needs_review_by_reason"].items()]
    lines += [
        f"Duplicate IDs:      {len(report['duplicate_source_ids'])}",
        f"Missing lat/lng:    {len(report['rows_missing_lat_lng'])}",
        f"Missing zip:        {len(report['rows_missing_zip'])}",
        "",
    ]
    if report["plz_check_applied"]:
        max_dist = report["max_plz_distance_km"]
        lines += [
            f"PLZ checks (GeoNames, limit {enrich.PLZ_DISTANCE_LIMIT_KM} km):",
            f"  Max distance to PLZ centroid: {max_dist:.1f} km" if max_dist is not None else "  Max distance to PLZ centroid: n/a",
            f"  {enrich.REASON_FAR_FROM_PLZ}:  {len(report['rows_far_from_plz'])}",
        ]
        lines += [f"    - {r['source_id']}  PLZ {r['zip']}  {r['plz_distance_km']} km" for r in report["rows_far_from_plz"]]
        lines += [f"  {enrich.REASON_UNKNOWN_PLZ}:              {len(report['rows_unknown_plz'])}"]
        lines += [f"    - {r['source_id']}  PLZ {r['zip']}" for r in report["rows_unknown_plz"]]
    else:
        lines += ["PLZ checks: not run (apply core.enrich.enrich first)"]
    lines += ["", "Fill rate (%):"]
    lines += [f"  {k:<20}{v:>6.1f}" for k, v in report["fill_rate_pct"].items()]
    lines += [
        "",
        'ID prefix check (report only; E -> "1...", G -> "00..."):',
        f"  Fits pattern:     {report['id_prefix_match_count']} / {report['rows_parsed']}",
        f"  Breaks pattern:   {len(report['id_prefix_breaks'])}",
    ]
    lines += [
        f"    - {r['source_id']}  herkunft:{r['herkunft']}  {r['name']}" for r in report["id_prefix_breaks"]
    ]
    lines += ["", "Warnings:"]
    lines += [f"  ! {w}" for w in report["warnings"]] or ["  none"]
    return "\n".join(lines)


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------

# Stored as text in XLSX so Excel keeps leading zeros
XLSX_TEXT_COLUMNS = ["source_id", "location_id", "zip"]
XLSX_MAX_COL_WIDTH = 60


def safe_path(path: str | Path, renamed: list | None = None) -> Path:
    """Never overwrite: if path exists, return name_2, name_3, ... and log the rename.

    Renames are also appended to `renamed` as (intended, actual) if a list is given.
    """
    path = Path(path)
    if not path.exists():
        return path
    n = 2
    while (candidate := path.with_name(f"{path.stem}_{n}{path.suffix}")).exists():
        n += 1
    logger.warning("%s already exists, writing %s instead", path, candidate.name)
    if renamed is not None:
        renamed.append((str(path), str(candidate)))
    return candidate


def _write_xlsx(frame: pd.DataFrame, path: Path, sheet_name: str) -> None:
    """Text columns as text, numbers as numbers, header frozen, column widths set."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]
        ws.freeze_panes = "A2"
        for idx, col in enumerate(frame.columns, start=1):
            letter = get_column_letter(idx)
            if col in XLSX_TEXT_COLUMNS:
                for cell in ws[letter][1:]:
                    cell.number_format = "@"
                    if cell.value is not None:
                        cell.value = str(cell.value)
            width = max([len(str(col))] + [len(str(v)) for v in frame[col].dropna()])
            ws.column_dimensions[letter].width = min(width + 2, XLSX_MAX_COL_WIDTH)


def export(
    df: pd.DataFrame, out_dir: str | Path, run_date: str | None = None, renamed: list | None = None, tag: str | None = None
) -> dict:
    """Write all rows and in-scope rows to CSV (utf-8-sig, ';', '.' decimal) and XLSX.

    Existing files are never overwritten (see safe_path). tag, if given, is appended to the
    file names (gtue_all_<date>_<tag>.csv). Returns {name: (path, rows)}.
    """
    run_date = run_date or date.today().isoformat()
    suffix = f"_{tag}" if tag else ""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    frames = {"all": df, "pruefstellen": df[df["in_scope"]]}
    result = {}
    for key, frame in frames.items():
        csv_path = safe_path(out_dir / f"gtue_{key}_{run_date}{suffix}.csv", renamed)
        frame.to_csv(csv_path, sep=";", encoding="utf-8-sig", index=False)
        result[f"{key}_csv"] = (csv_path, len(frame))

        xlsx_path = safe_path(out_dir / f"gtue_{key}_{run_date}{suffix}.xlsx", renamed)
        _write_xlsx(frame, xlsx_path, sheet_name=key)
        result[f"{key}_xlsx"] = (xlsx_path, len(frame))
    return result


def write_recon_run(
    df: pd.DataFrame,
    source_file: str | Path,
    root: str | Path | None = None,
    run_date: str | None = None,
    started_at: datetime | None = None,
) -> tuple[Path, dict]:
    """Mode "recon": export an offline-parsed sample response into a run folder of its own in test_output/.

    No requests. Writes the four exports plus run_summary.json/.md; returns (run folder, export result).
    """
    root = Path(root) if root is not None else paths.project_root()
    started = started_at or datetime.now()
    run_date = run_date or started.date().isoformat()
    run_folder = paths.new_run_dir(SLUG, "recon", started, root)
    renamed: list = []
    written = export(df, run_folder, run_date, renamed)
    summary = {
        "source": SOURCE_ORG,
        "mode": "recon",
        "run_date": run_date,
        "started_at": started.isoformat(timespec="seconds"),
        "source_file": _relative(Path(source_file), root),
        "rows": len(df),
        "in_scope": int(df["in_scope"].sum()),
        "files": {"run_folder": _relative(run_folder, root)}
        | {key: _relative(path, root) for key, (path, _rows) in written.items()},
    }
    (run_folder / paths.RUN_SUMMARY_JSON).write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    text = "\n".join(
        [f"GTÜ recon export (no network): {summary['rows']} rows, {summary['in_scope']} in scope",
         f"Source file: {summary['source_file']}", "", "Files:"]
        + [f"  {k:<18}{v}" for k, v in summary["files"].items()]
    )
    (run_folder / paths.RUN_SUMMARY_MD).write_text(render_summary_md(run_folder.name, text), encoding="utf-8")
    return run_folder, written


# --------------------------------------------------------------------------
# Fetch
# --------------------------------------------------------------------------

SLUG = "gtue"
PROJECT_ROOT = paths.project_root()

API_URL = "https://gtue-finder-bff.smartmaps.app/addresssearch/ByBranchesAndAddress"
FINDER_ORIGIN = "https://gtue-finder.smartmaps.app"
API_KEY_ENV = "GTUE_API_KEY"
TOP = 20  # server cap

PACE_S = 3.0  # between requests, plus 0-JITTER_S random
JITTER_S = 1.0
TIMEOUT_S = 30
MAX_ATTEMPTS = 3  # timeouts / connection errors / 5xx
BACKOFF_S = 5.0  # 5 s, then 10 s
RATE_LIMIT_WAIT_S = 60  # on 429, once

# Reference query: the one national query, made at the start (and, in a full run, at the end).
# The server caps MaxRadius at 500 km (confirmed 2026-09-29), so send exactly 500 km. The expected actual
# centre (SearchData.Location) is 51.15844 N, 10.44595 E; the German border is at most 454 km from it
# (List/Sylt), so the 500 km circle covers all of Germany and its MaxCount is the number of registered
# locations inside it. The coverage proof depends on that centre: a run stops if it comes back more
# than REFERENCE_CENTRE_TOLERANCE_KM away.
REFERENCE_ZIP = "99986"
REFERENCE_CITY = "Niederdorla"
REFERENCE_RADIUS_KM = 500  # = SERVER_RADIUS_CAP_KM
REFERENCE_CENTRE = (51.15844, 10.44595)  # (lat, lng)
REFERENCE_CENTRE_TOLERANCE_KM = 5
REFERENCE_MAX_BORDER_KM = 454  # farthest German border point from the reference centre
SERVER_RADIUS_CAP_KM = 500
# Cap probe (early stop only): MaxCount for München 80686 at two radii just below the cap. Equal counts
# mean the cap has dropped below what covers Germany from the reference centre.
CAP_PROBE_ZIP = "80686"
CAP_PROBE_CITY = "München"
CAP_PROBE_RADII_KM = (455, 460)
SAMPLE_GRID_SQUARES = 2
# sample run: the reference query, the cap probe (2 requests) and 2 grid squares
SAMPLE_MAX_REQUESTS = 1 + len(CAP_PROBE_RADII_KM) + SAMPLE_GRID_SQUARES
FULL_MAX_REQUESTS = 2000  # safety ceiling (live requests). Run 2026-09-29 hit it without re-query/early stop;
# offline replay of that run: re-query + early stop ≈ 900 requests, re-query alone 1,750–1,950

FULL_RUN_NOT_CONFIRMED = (
    "Full run not confirmed: call run(mode='full', confirm_full_run=True). "
    "Run it only after a successful sample run (mode='sample')."
)

GEOCODING_OUTLIER_KM = 5  # server centre this far from the queried PLZ centroid is listed
MAX_GERMAN_DROP = 0.10  # alert if German locations drop more than this vs the last full run
SITE_TYPES = ("pruefstelle", "partner_office", "unknown")


def _tier_counts(frame: pd.DataFrame) -> dict[str, int]:
    """Rows per scope tier (all tiers listed, in the order of SCOPE_TIERS)."""
    return {t: int(frame["scope_tier"].eq(t).sum()) for t in SCOPE_TIERS}


class KeyRejectedError(StopRun):
    """HTTP 401: the finder key no longer works."""


class ApiError(StopRun):
    """The API failed or answered in an unexpected way; stop rather than guess."""


class RequestLimitReached(StopRun):
    """The client's hard request limit was reached."""


def load_api_key(root: str | Path = PROJECT_ROOT) -> str:
    """GTUE_API_KEY from .env (falls back to the environment). Never print, log or save it."""
    key = dotenv_values(Path(root) / ".env").get(API_KEY_ENV) or os.environ.get(API_KEY_ENV)
    if not key or not key.strip():
        raise StopRun(f"{API_KEY_ENV} is not set: store it with tools\\store-secret.ps1 and run via tools\\with-secrets.ps1")
    return key.strip()


class GtueClient:
    """Paced, retrying client for the finder API. One request at a time, never parallel."""

    def __init__(
        self,
        api_key: str,
        max_requests: int | None = None,
        session: requests.Session | None = None,
        pace_s: float = PACE_S,
        jitter_s: float = JITTER_S,
        timeout_s: float = TIMEOUT_S,
        sleep=time.sleep,
        clock=time.monotonic,
    ):
        self._session = session if session is not None else requests.Session()
        # User-Agent stays requests' default: we don't impersonate a browser
        self._headers = {
            "x-api-key": api_key,
            "Origin": FINDER_ORIGIN,
            "Referer": FINDER_ORIGIN + "/",
            "Accept": "application/json",
        }
        self.max_requests = max_requests
        self.pace_s = pace_s
        self.jitter_s = jitter_s
        self.timeout_s = timeout_s
        self._sleep = sleep
        self._clock = clock
        self._last_request_at = None
        self.requests_made = 0
        self.last_http_status = None

    def __repr__(self) -> str:  # never expose the key
        return f"GtueClient(requests_made={self.requests_made}, max_requests={self.max_requests})"

    def _send(self, params: dict) -> requests.Response:
        if self.max_requests is not None and self.requests_made >= self.max_requests:
            raise RequestLimitReached(
                f"Request limit reached: {self.requests_made} live requests sent (max_requests={self.max_requests})"
            )
        if self._last_request_at is not None:
            wait = self.pace_s + random.uniform(0, self.jitter_s) - (self._clock() - self._last_request_at)
            if wait > 0:
                self._sleep(wait)
        self.requests_made += 1
        try:
            return self._session.get(API_URL, params=params, headers=self._headers, timeout=self.timeout_s)
        finally:
            self._last_request_at = self._clock()

    def query(self, zip: str, city: str, radius_m: int) -> dict:
        """One address search around a PLZ. Returns the raw JSON response."""
        params = {
            "Top": TOP,
            "MaxRadius": int(radius_m),
            "Branches": "GTUE",
            "City": city or " ",
            "Zip": zip,
            "Street": " ",
            "IsoCountryCode": "DE",
        }
        attempt = 0
        rate_limited = False
        while True:
            attempt += 1
            try:
                response = self._send(params)
            except (requests.Timeout, requests.ConnectionError) as exc:
                self.last_http_status = None
                if attempt >= MAX_ATTEMPTS:
                    raise ApiError(f"GTÜ API unreachable after {attempt} attempts (PLZ {zip}): {type(exc).__name__}") from exc
                self._sleep(BACKOFF_S * 2 ** (attempt - 1))
                continue

            status = response.status_code
            self.last_http_status = status
            if status == 200:
                try:
                    return response.json()
                except ValueError as exc:
                    raise ApiError(f"GTÜ API returned invalid JSON (PLZ {zip})") from exc
            if status == 401:
                raise KeyRejectedError("GTÜ API returned 401: key rejected, probably rotated, update .env")
            if status == 429:
                if rate_limited:
                    raise ApiError(f"GTÜ API rate limit (429) again after waiting {RATE_LIMIT_WAIT_S} s: stopping the run")
                rate_limited = True
                self._sleep(RATE_LIMIT_WAIT_S)
                continue
            if 500 <= status < 600 and attempt < MAX_ATTEMPTS:
                self._sleep(BACKOFF_S * 2 ** (attempt - 1))
                continue
            raise ApiError(f"GTÜ API returned HTTP {status} (PLZ {zip}, attempt {attempt})")


RAW_REF_SEP = "#"
_RAW_LINE_FILE = re.compile(rb'^\{"file": "([^"]+)"')


def raw_ref(run_folder: Path, root: Path, name: str) -> str:
    """The query log's raw_file value: '<run folder rel. to root>/raw.jsonl#<name>'."""
    return f"{_relative(Path(run_folder) / paths.RAW_FILE, root)}{RAW_REF_SEP}{name}"


class RawReader:
    """Random access to raw.jsonl files, which can be hundreds of MB: the first access to a file scans it
    once and remembers where each line starts, so a lookup parses one line only."""

    def __init__(self):
        self._index: dict = {}  # resolved path -> {file name: (offset, length)}

    def _scan(self, path: Path) -> dict:
        index, offset = {}, 0
        with path.open("rb") as handle:
            for line in handle:
                found = _RAW_LINE_FILE.match(line)
                if found:
                    index[found.group(1).decode("utf-8")] = (offset, len(line))
                offset += len(line)
        self._index[path] = index
        return index

    def register(self, path: Path, name: str, span: tuple[int, int]) -> None:
        """A line the running run has just written: known without re-scanning the growing file."""
        path = Path(path).resolve()
        index = self._index.get(path)
        if index is None:
            index = self._scan(path)
        index[name] = span

    def get(self, path: Path, name: str) -> dict:
        """The response stored under `name`. Raises OSError/ValueError/KeyError if it can't be read."""
        path = Path(path).resolve()
        index = self._index.get(path)
        if index is None:
            index = self._scan(path)
        offset, length = index[name]
        with path.open("rb") as handle:
            handle.seek(offset)
            record = json.loads(handle.read(length).decode("utf-8"))
        return record["response"]


def iter_raw(path: str | Path):
    """(file name, response) for every line of a raw.jsonl, in request order. Unreadable lines yield (None, None)."""
    with Path(path).open("rb") as handle:
        for line in handle:
            if not line.strip():
                continue
            try:
                record = json.loads(line.decode("utf-8"))
                yield record["file"], record["response"]
            except (ValueError, KeyError, TypeError):
                yield None, None


class RawStore:
    """Appends every raw response of a run to <run folder>/raw.jsonl, one line per request, in request order:
    {"file": "<nnnn>.json", "response": <the response as received>}. A second store on the same folder
    continues the numbering and never rewrites earlier lines. save() returns the reference that goes into the
    query log's raw_file column (see raw_ref)."""

    def __init__(self, directory: str | Path, root: str | Path):
        self.directory = Path(directory)
        self.root = Path(root)
        self.directory.mkdir(parents=True, exist_ok=True)
        self.path = self.directory / paths.RAW_FILE
        self._n = 0
        if self.path.exists():
            with self.path.open("rb") as handle:
                self._n = sum(1 for line in handle if line.strip())
        self.last_span: tuple[int, int] | None = None  # (offset, length) of the line saved last

    def save(self, data: dict) -> str:
        self._n += 1
        name = f"{self._n:04d}.json"
        line = (json.dumps({"file": name, "response": data}, ensure_ascii=False) + "\n").encode("utf-8")
        offset = self.path.stat().st_size if self.path.exists() else 0
        with self.path.open("ab") as handle:
            handle.write(line)
        self.last_span = (offset, len(line))
        return raw_ref(self.directory, self.root, name)


def to_query_result(
    data: dict, http_status: int | None = None, raw_file: str | None = None, cached: bool = False
) -> QueryResult:
    """Map a finder response to the grid's QueryResult. Unexpected shapes stop the run."""
    if data.get("Exceptions") or data.get("InternalError"):
        raise ApiError(f"GTÜ API reported an error: {data.get('Exceptions') or data.get('InternalError')}")
    try:
        paging = data["Paging"]
        location = data["SearchData"]["Location"]
        count = int(paging["Count"])
        max_count = int(paging["MaxCount"])
        center_lat = float(location["YCoord"])  # Y = latitude
        center_lng = float(location["XCoord"])  # X = longitude
        address_items = list(data["AddressItems"])
    except (KeyError, TypeError, ValueError) as exc:
        raise ApiError(f"Unexpected GTÜ response format ({type(exc).__name__}: {exc})") from exc

    items = {}
    for item in address_items:
        pfk = _clean(((item.get("BasicData") or {}).get("Identifiers") or {}).get("ProviderForeignKey"))
        if pfk is None:
            raise ApiError("GTÜ response contains an entry without ProviderForeignKey")
        items.setdefault(pfk, item)
    # If fewer entries arrive than Count claims, the missing ones count as not returned
    return QueryResult(
        min(count, len(address_items)), max_count, center_lat, center_lng, items, http_status, raw_file, cached
    )


class ResponseCache:
    """Same-day response cache: (zip, city, radius in m) -> raw response.

    A cache hit sends no request (and doesn't count against the request limit), so a re-run
    on the same date, or a repeated identical query within a run, is answered from disk.
    Filled from earlier runs of the same date in the same area (output/ or test_output/): their
    query_log.csv says which query each raw.jsonl line answers. Only that date's responses are used:
    the finder's data can change from one day to the next.
    hits_by_run says which run folder each answer came from (the running run's own folder counts too).
    """

    def __init__(self, root: str | Path):
        self.root = Path(root)
        self._refs: dict = {}  # key -> raw_file reference ('<run folder>/raw.jsonl#<name>')
        self._memory: dict = {}  # key -> response, when there is no raw file (simulated runs)
        self._reader = RawReader()
        self.loaded = 0  # entries taken from earlier runs
        self.loaded_from: dict = {}  # run folder name -> entries added
        self.hits = 0
        self.hits_by_run: dict = {}  # run folder name -> answers served from it

    @staticmethod
    def key(zip: str, city: str | None, radius_m: int) -> tuple:
        return (str(zip), city or "", int(radius_m))

    def load_runs(self, area_dir: str | Path, run_date: str, exclude: str | Path | None = None) -> int:
        """Add every successful query of the earlier runs of run_date in area_dir (each needs query_log.csv and
        raw.jsonl; the folder `exclude` is the running run). Returns entries added."""
        area = Path(area_dir).resolve()
        exclude = Path(exclude).resolve() if exclude is not None else None
        added = 0
        for folder in sorted(p for p in area.glob(f"{run_date}_*") if p.is_dir()):
            if folder.resolve() == exclude:
                continue
            log_path = folder / paths.QUERY_LOG_FILE
            if not log_path.exists() or not (folder / paths.RAW_FILE).exists():
                continue
            log = pd.read_csv(log_path, sep=";", encoding="utf-8-sig", dtype=str, keep_default_na=False)
            if not {"zip", "city", "radius_m", "http_status", "raw_file"} <= set(log.columns):
                continue
            for row in log.to_dict("records"):
                ref = row["raw_file"]
                if row["http_status"] not in ("200", "200.0") or RAW_REF_SEP not in ref or not row["zip"]:
                    continue
                target = (self.root / ref.rsplit(RAW_REF_SEP, 1)[0]).resolve()
                # only raw.jsonl files of the same day in this area
                if target.name != paths.RAW_FILE or target.parent.parent != area or not target.parent.name.startswith(f"{run_date}_"):
                    continue
                key = self.key(row["zip"], row["city"], int(float(row["radius_m"])))
                if key not in self._refs:
                    self._refs[key] = ref
                    added += 1
                    self.loaded_from[folder.name] = self.loaded_from.get(folder.name, 0) + 1
        self.loaded += added
        return added

    def _read(self, ref: str) -> dict:
        rel, name = ref.rsplit(RAW_REF_SEP, 1)
        return self._reader.get(self.root / rel, name)

    def get(self, key: tuple) -> tuple[dict, str | None] | None:
        """(response, raw file reference) or None. An unreadable raw file counts as a miss."""
        if key in self._memory:
            found = (self._memory[key], None)
        elif key in self._refs:
            try:
                found = (self._read(self._refs[key]), self._refs[key])
            except (OSError, ValueError, KeyError):
                del self._refs[key]
                return None
        else:
            return None
        self.hits += 1
        if found[1] is not None:
            run = Path(found[1].rsplit(RAW_REF_SEP, 1)[0]).parent.name
            self.hits_by_run[run] = self.hits_by_run.get(run, 0) + 1
        return found

    def put(self, key: tuple, raw_file: str | None, data: dict, span: tuple[int, int] | None = None) -> None:
        if raw_file:
            self._refs[key] = raw_file
            if span is not None:  # the line the running run has just written
                rel, name = raw_file.rsplit(RAW_REF_SEP, 1)
                self._reader.register(self.root / rel, name, span)
        else:
            self._memory[key] = data


def make_query_fn(api, raw_store: RawStore | None = None, cache: ResponseCache | None = None):
    """Grid query function around anything with query(zip, city, radius_m) -> dict.
    With a cache, a query already answered today is served from it instead of the API."""

    def query_fn(point: QueryPoint, radius_km: int) -> QueryResult:
        key = ResponseCache.key(point.zip, point.city, radius_km * 1000)
        if cache is not None and (hit := cache.get(key)) is not None:
            data, raw_file = hit
            return to_query_result(data, 200, raw_file, cached=True)
        data = api.query(point.zip, point.city, radius_km * 1000)
        raw_file = raw_store.save(data) if raw_store else None
        result = to_query_result(data, api.last_http_status, raw_file)
        if cache is not None:
            cache.put(key, raw_file, data, span=raw_store.last_span if raw_store else None)  # only responses that parsed
        return result

    return query_fn


# --------------------------------------------------------------------------
# Simulator (offline stand-in for the API)
# --------------------------------------------------------------------------

def simulated_locations(
    centroids: pd.DataFrame, n_total: int = 1200, seed: int = 0, clusters: dict = BIG_CITIES
) -> pd.DataFrame:
    """Fake locations (by default with dense clusters around 5 big cities), GTÜ-style IDs and Herkunft."""
    locations = fake_locations(centroids, n_total=n_total, clusters=clusters, seed=seed)
    rng = random.Random(seed)
    herkunft = ["E" if rng.random() < 0.65 else "G" for _ in range(len(locations))]
    locations["herkunft"] = herkunft
    locations["id"] = [f"1{i:05d}000" if h == "E" else f"00{i:07d}" for i, h in enumerate(herkunft)]
    return locations


def _fake_item(row) -> dict:
    branches = [{"BranchCode": "GTUE", "BranchText": "GTÜ-Ingenieurbüro"}]
    if row.herkunft == "E":
        branches.append({"BranchCode": VOR_ORT_CODE, "BranchText": "GTÜ-Ingenieurbüro mit amtlicher Fahrzeugprüfung vor Ort"})
    return {
        "BasicData": {
            "Identifiers": {"ProviderForeignKey": row.id},
            "Address": {
                "CompanyName": f"SIM Prüfstelle {row.id}",
                "Street": "Simulierte Straße 1",
                "Zip": row.zip,
                "City": row.city,
                "Country": "D",
            },
            "Geo": {"XCoord": f"{row.lng:.5f}", "YCoord": f"{row.lat:.5f}"},
            "Contact": {},
            "BusinessData": {},
            "MemoItems": [
                {"Type": "JSON_INFO", "Text": json.dumps({"HiddenPoiInfos": {"Herkunft": row.herkunft, "Stand": "2026-09-20"}})}
            ],
            "BranchListElements": branches,
            "ObjectListItems": [{"Headline": HU_HEADLINE}],
        }
    }


class FakeGtueApi:
    """Offline stand-in for GtueClient: same query() signature, GTÜ-shaped responses.
    The centre it reports is offset from the PLZ centroid, like the real geocoder."""

    def __init__(self, locations: pd.DataFrame, centroids: pd.DataFrame, seed: int = 0, center_offset_km: float = 2.0):
        self.finder = SimulatedFinder(locations, top=TOP, center_offset_km=center_offset_km, seed=seed)
        self.items = {row.id: _fake_item(row) for row in locations.itertuples(index=False)}
        self._points = centroids.set_index("zip")
        self.requests_made = 0
        self.last_http_status = None

    def query(self, zip: str, city: str, radius_m: int) -> dict:
        self.requests_made += 1
        self.last_http_status = 200
        row = self._points.loc[zip]
        point = QueryPoint(zip, city, float(row["lat"]), float(row["lng"]))
        c_lat, c_lng, max_count, ids = self.finder.search(point, radius_m / 1000)
        return {
            "Paging": {"Page": 1, "MaxPage": 1, "Count": len(ids), "MaxCount": max_count},
            "SearchData": {"Location": {"XCoord": f"{c_lng:.5f}", "YCoord": f"{c_lat:.5f}"}},
            "AddressItems": [self.items[i] for i in ids],
            "Exceptions": [],
            "InternalError": None,
        }


# --------------------------------------------------------------------------
# Run
# --------------------------------------------------------------------------

MODES = ("simulated", "sample", "full")
OPEN_SQUARE_COLUMNS = ["kind", "square_km", "lat_min", "lat_max", "lng_min", "lng_max", "zip", "reason"]


@dataclass
class RunResult:
    summary: dict
    df: pd.DataFrame
    query_log: pd.DataFrame


def _relative(path: Path, root: Path) -> str:
    try:
        return Path(path).resolve().relative_to(Path(root).resolve()).as_posix()
    except ValueError:
        return str(path)


def run_paths(run_folder: str | Path) -> dict:
    """File locations inside one run folder (the folder itself comes from core.paths.new_run_dir)."""
    run_folder = Path(run_folder)
    return {
        "folder": run_folder,
        "query_log": run_folder / paths.QUERY_LOG_FILE,
        "summary": run_folder / paths.RUN_SUMMARY_JSON,
        "summary_md": run_folder / paths.RUN_SUMMARY_MD,
        "open_squares": run_folder / paths.OPEN_SQUARES_FILE,
    }


def _query_point(centroids: pd.DataFrame, zip: str) -> QueryPoint:
    row = centroids.loc[centroids["zip"] == zip].iloc[0]
    return QueryPoint(zip, str(row["city"]), float(row["lat"]), float(row["lng"]))


def _countries(items: dict) -> dict:
    """Address.Country values of raw items (id -> item)."""
    counts: dict[str, int] = {}
    for item in items.values():
        country = _clean(((item.get("BasicData") or {}).get("Address") or {}).get("Country")) or "missing"
        counts[country] = counts.get(country, 0) + 1
    return counts


def _baseline(area_dir: str | Path, exclude: str | Path | None = None) -> dict | None:
    """Reference for the drop check: the most recent run in output/ whose run_summary.json says PASS.
    Failed or incomplete runs never count. The run must also have collected German locations and must not
    have raised the drop alert itself, so one bad run can't become the new normal. None: no baseline."""
    best = None
    exclude = Path(exclude).resolve() if exclude is not None else None
    for path in sorted(Path(area_dir).glob(f"*_full*/{paths.RUN_SUMMARY_JSON}")):
        if path.parent.resolve() == exclude:
            continue
        try:
            s = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if s.get("mode") != "full" or s.get("status") != "PASS" or s.get("aborted") or s.get("german_drop_alert"):
            continue
        if not s.get("unique_german_locations"):  # None or 0: nothing to compare with
            continue
        if best is None or s["started_at"] > best["started_at"]:
            best = {**s, "_file": path.parent.name}
    return best


def _join(counts: dict) -> str:
    return ", ".join(f"{k}: {v}" for k, v in counts.items())


def _coords(item: dict) -> tuple[float, float] | None:
    geo = (item.get("BasicData") or {}).get("Geo") or {}
    lat, lng = _to_float(geo.get("YCoord")), _to_float(geo.get("XCoord"))
    return (lat, lng) if lat is not None and lng is not None else None


def _country(item: dict) -> str | None:
    return _clean(((item.get("BasicData") or {}).get("Address") or {}).get("Country"))


def _in_circle(centre: tuple[float, float], coords: tuple[float, float], radius_km: float) -> bool:
    return float(haversine_km(centre[0], centre[1], coords[0], coords[1])) <= radius_km


def circle_split(items: dict, centre: tuple[float, float], radius_km: float = REFERENCE_RADIUS_KM):
    """Collected raw items (id -> item) split by great-circle distance from centre:
    (ids inside the circle, ids outside, entries without coordinates)."""
    inside, outside, no_coords = set(), set(), 0
    for key, item in items.items():
        coords = _coords(item)
        if coords is None:
            no_coords += 1
        elif _in_circle(centre, coords, radius_km):
            inside.add(key)
        else:
            outside.add(key)
    return inside, outside, no_coords


def count_check(items: dict, centre, observations: list[int], radius_km: float = REFERENCE_RADIUS_KM) -> dict:
    """Criterion (b): N_circle (collected unique IDs within radius_km of the reference's actual centre)
    against the reference MaxCount. observations = every MaxCount the reference returned during the
    run (start, early-stop rechecks, end); if they differ, N_circle may be anywhere in [min, max].

    verdict: "match" | "gap" (N_circle below the MaxCount) | "excess" (above: data change, or the
    server's count or cap changed: investigate). Also counts rows with Country != "D" (missing counts
    as abroad) inside and outside the circle, and the German total (rows with Country == "D")."""
    if centre is None or not observations:
        return {"evaluated": False, "problems": []}
    inside, outside, no_coords = circle_split(items, centre, radius_km)
    low, high = min(observations), max(observations)
    n = len(inside)
    german = {k for k, item in items.items() if _country(item) == HOME_COUNTRY}
    abroad_inside = len(inside - german)
    verdict = "match" if low <= n <= high else ("gap" if n < low else "excess")
    return {
        "evaluated": True,
        "radius_km": radius_km,
        "centre": list(centre),
        "n_circle": n,
        "observations": list(observations),
        "reference_low": low,
        "reference_high": high,
        "verdict": verdict,
        "gap": low - n if verdict == "gap" else None,
        "excess": n - high if verdict == "excess" else None,
        "abroad_inside": abroad_inside,
        "abroad_outside": len(outside - german),
        "german_total": len(german),
        "german_outside_circle": len(german & outside),
        "items_without_coordinates": no_coords,
        "problems": [],
    }


class ReferenceCentreError(StopRun):
    """The reference query's actual centre is too far from the expected one: the coverage proof
    (a circle that covers Germany) no longer holds."""


def reference_offset_km(res: QueryResult) -> float:
    """Distance between the reference's actual centre and the expected one."""
    return round(float(haversine_km(res.center_lat, res.center_lng, *REFERENCE_CENTRE)), 2)


def check_reference_centre(res: QueryResult) -> float:
    """Stop the run if the reference's actual centre is more than REFERENCE_CENTRE_TOLERANCE_KM from
    the expected one. Returns the offset."""
    offset = reference_offset_km(res)
    if offset > REFERENCE_CENTRE_TOLERANCE_KM:
        raise ReferenceCentreError(
            f"Reference query ({REFERENCE_ZIP} {REFERENCE_CITY}, {REFERENCE_RADIUS_KM} km) came back centred on "
            f"{res.center_lat}, {res.center_lng}: {offset} km from the expected {REFERENCE_CENTRE[0]}, "
            f"{REFERENCE_CENTRE[1]} (limit {REFERENCE_CENTRE_TOLERANCE_KM} km). The coverage proof depends on this "
            "centre, so the run stops. Check the geocoding of the reference address before running again."
        )
    return offset


class EarlyStop:
    """Stopping rule E: grid.run_grid(stop_when=...) stops once N_circle reaches the reference MaxCount.

    N_circle = collected unique IDs within REFERENCE_RADIUS_KM of the reference's actual centre, never
    all collected IDs (entries abroad beyond the circle must not trigger the stop). When N_circle
    reaches the start MaxCount the reference is queried once more (live) and the run stops only if
    N_circle still equals the MaxCount; if the reference changed, the run continues against the new
    value. N_circle above the MaxCount also stops the run (overshoot), and the run then FAILs
    (criterion (b): investigate). A failed recheck, or a reference centre that moved, switches the
    early stop off for the rest of the run (`error`)."""

    def __init__(self, live_fn, point: QueryPoint, start: QueryResult, radius_km: int = REFERENCE_RADIUS_KM):
        self.live_fn, self.point, self.radius_km = live_fn, point, radius_km
        self.centre = (start.center_lat, start.center_lng)
        self.target = start.max_count
        self.target_changes: list[tuple[int, int]] = []
        self.rechecks = 0
        self.triggered = False
        self.overshoot = False
        self.error: str | None = None
        self._done = 0
        self._inside = 0

    def n_circle(self, result: GridResult) -> int:
        """Incremental: only items collected since the last call are measured."""
        for _key, item in islice(result.items.items(), self._done, None):
            self._done += 1
            coords = _coords(item)
            if coords is not None and _in_circle(self.centre, coords, self.radius_km):
                self._inside += 1
        return self._inside

    def __call__(self, result: GridResult) -> bool:
        if self.error:
            return False
        n = self.n_circle(result)
        if n < self.target:
            return False
        res = reference_recheck(self.live_fn, self.point, self.radius_km, result)
        if isinstance(res, dict):
            self.error = f"recheck failed: {res['error']}"
            return False
        self.rechecks += 1
        offset = reference_offset_km(res)
        if offset > REFERENCE_CENTRE_TOLERANCE_KM:
            self.error = f"reference centre moved {offset} km during the run"
            return False
        if res.max_count != self.target:
            self.target_changes.append((self.target, res.max_count))
            self.target = res.max_count
        if n == self.target:
            self.triggered = True
            return True
        if n > self.target:
            self.overshoot = True
            return True
        return False  # the reference grew: continue against the new value

    def summary(self) -> dict:
        return {
            "triggered": self.triggered,
            "overshoot": self.overshoot,
            "rechecks": self.rechecks,
            "target_changes": [list(c) for c in self.target_changes],
            "error": self.error,
        }


def run_cap_probe(live_fn, point: QueryPoint, result: GridResult) -> dict:
    """MaxCount for Zip 80686 / München at 455 and at 460 km (2 live requests). Equal counts mean the
    server's radius cap has dropped below what covers Germany from the reference centre."""
    counts = [cap_probe(live_fn, point, km, result).max_count for km in CAP_PROBE_RADII_KM]
    return {"radii_km": list(CAP_PROBE_RADII_KM), "max_counts": counts, "equal": counts[0] == counts[1]}


def run_status(
    mode: str, grid_result: GridResult, abort: BaseException | None = None, count: dict | None = None
) -> tuple[str, list[str]]:
    """Rule C4. PASS needs both:
    (a) every grid square is complete (Count == MaxCount and the radius covers the square from the
        server's actual centre, checked in grid.run_grid): 0 coverage warnings, 0 squares not
        queried, 0 errored queries. An early stop that triggered (EarlyStop) replaces the warnings
        and not-queried part of (a); errors still fail;
    (b) N_circle == the reference MaxCount (see count_check; `count` is its result). A gap or an
        excess is FAIL. Not evaluated after an abort (already FAIL) or in sample mode (a sample
        never collects everything).
    Not used: the drop check vs the last full run (an alert of its own)."""
    reasons = []
    if abort is not None:
        reasons.append(f"run aborted: {type(abort).__name__}: {abort}")
    if mode == "sample":
        reasons.append(
            f"sample mode: reference query + cap probe + {SAMPLE_GRID_SQUARES} grid squares only, not a full crawl"
        )
    early = grid_result.stopped_early
    n_warn = 0 if early else len(grid_result.coverage_warnings)
    n_not_queried = 0 if early else grid_result.squares_not_queried
    n_errors = grid_result.count("error")
    if n_warn:
        reasons.append(f"{n_warn} squares still incomplete at minimum size (coverage warnings)")
    if n_not_queried:
        reasons.append(f"{n_not_queried} squares not queried")
    if n_errors:
        reasons.append(f"{n_errors} queries errored")
    if count and count.get("evaluated") and abort is None and mode != "sample":
        if count["verdict"] == "gap":
            reasons.append(
                f"N_circle {count['n_circle']} < reference MaxCount {count['reference_low']} (gap {count['gap']})"
            )
        elif count["verdict"] == "excess":
            reasons.append(
                f"N_circle {count['n_circle']} > reference MaxCount {count['reference_high']}: investigate"
                " (data change, or the server's count or cap changed)"
            )
    if abort is None and count:
        reasons.extend(count.get("problems", []))
    if not reasons and not any(r["status"] in ("complete", "split", "warning") for r in grid_result.log):
        reasons.append("no grid square was queried")  # never PASS an empty crawl
    return ("FAIL", reasons) if reasons else ("PASS", [])


def drop_check(german: int | None, baseline: dict | None) -> tuple[float | None, str | None]:
    """(change in %, alert text or None) for German locations vs the baseline full run."""
    if not baseline or german is None or not baseline["unique_german_locations"]:
        return None, None
    before = baseline["unique_german_locations"]
    change_pct = round((german - before) / before * 100, 1)
    if german < before * (1 - MAX_GERMAN_DROP):
        return change_pct, (
            f"German locations dropped {-change_pct}% vs the last full run "
            f"({baseline['run_date']}: {before} -> {german}; alert above {MAX_GERMAN_DROP:.0%})"
        )
    return change_pct, None


def _completion(grid_result: GridResult, abort: BaseException | None, early: dict) -> str:
    if abort is not None:
        return "aborted"
    if early.get("overshoot"):
        return "stopped early: N_circle above the reference MaxCount (investigate)"
    if grid_result.stopped_early:
        return "early stop: N_circle reached the reference MaxCount (open squares listed)"
    if grid_result.squares_not_queried or grid_result.coverage_warnings:
        return "grid incomplete"
    return "grid complete"


def outcome_label(
    status: str,
    grid_result: GridResult,
    abort: BaseException | None = None,
    count: dict | None = None,
    mode: str | None = None,
) -> str:
    """"PASS (grid complete + count match)", "PASS (early stop, count match)" or "FAIL (...)" with what
    is open. If entries with Country != "D" lie inside the reference circle, a PASS carries a note:
    the German total is then the number of rows with Country == "D"."""
    if status == "PASS":
        label = "PASS (early stop, count match)" if grid_result.stopped_early else "PASS (grid complete + count match)"
        if count and count.get("abroad_inside"):
            label += (
                f" with note: {count['abroad_inside']} rows with Country != D inside the circle,"
                f" German total {count['german_total']}"
            )
        return label
    warned = len(grid_result.coverage_warnings)
    pending = grid_result.squares_not_queried
    parts = []
    if abort is not None:
        parts.append("run aborted")
    if warned or pending:
        parts.append(f"{warned + pending} squares open: {warned} coverage warnings, {pending} not queried")
    if count and count.get("evaluated") and abort is None and mode != "sample" and count["verdict"] != "match":
        parts.append(
            f"count gap {count['gap']}" if count["verdict"] == "gap" else f"N_circle {count['excess']} above MaxCount"
        )
    return f"FAIL ({'; '.join(parts)})" if parts else "FAIL"


def _evaluate(mode, grid_result, abort, reference, reference_end) -> dict:
    """Criteria (a) and (b) of rule C4 for a finished (or aborted) run: the status and reasons, the
    count check and the reference MaxCounts seen (start, early-stop rechecks, end)."""
    log = grid_result.log
    recheck_rows = [r for r in log if r["status"] == "reference_recheck"]
    end_row = next((r for r in reversed(log) if r["status"] == "reference_end"), None)
    observations = (
        ([reference.max_count] if reference else [])
        + [int(r["max_count"]) for r in recheck_rows]
        + ([int(end_row["max_count"])] if end_row else [])
    )
    centre = (reference.center_lat, reference.center_lng) if reference else None
    count = count_check(grid_result.items, centre, observations)
    end_offset = None
    if isinstance(reference_end, QueryResult):
        end_offset = reference_offset_km(reference_end)
        if end_offset > REFERENCE_CENTRE_TOLERANCE_KM:
            count["problems"].append(
                f"reference centre moved {end_offset} km by the end of the run (limit {REFERENCE_CENTRE_TOLERANCE_KM} km)"
            )
    if count.get("german_outside_circle"):
        count["problems"].append(
            f"{count['german_outside_circle']} German rows lie outside the reference circle: it no longer covers Germany"
        )
    status, reasons = run_status(mode, grid_result, abort, count)
    return {
        "status": status, "reasons": reasons, "count": count, "observations": observations,
        "end_offset": end_offset, "recheck_rows": recheck_rows, "end_row": end_row,
    }


def _build_summary(
    mode, run_date, started, runtime_s, api, grid_result, reference, df, abort, files, renamed, baseline, cache=None,
    reference_end=None, early=None, max_requests=None, limit_reached=False,
) -> dict:
    """Run summary. status/status_reasons come from run_status() only (rule C4); the drop check is
    reported next to it and never changes it."""
    early = early or {}
    log = grid_result.log
    start_row = next((r for r in log if r["status"] == "reference_check"), None)
    ev = _evaluate(mode, grid_result, abort, reference, reference_end)
    status, reasons, count, observations = ev["status"], ev["reasons"], ev["count"], ev["observations"]
    recheck_rows, end_row, end_offset = ev["recheck_rows"], ev["end_row"], ev["end_offset"]
    centre = (reference.center_lat, reference.center_lng) if reference else None

    alerts = []
    if len(set(observations)) > 1:
        alerts.append(
            f"reference MaxCount changed during the run: {' -> '.join(str(o) for o in observations)}"
            f" (N_circle may be anywhere between {min(observations)} and {max(observations)})"
        )
    german_df = df[df["country"].eq(HOME_COUNTRY)] if df is not None else None
    german = len(german_df) if german_df is not None else None
    unique = len(grid_result.items)

    if mode != "full":
        drop_status = "n/a (only full runs are compared)"
        change_pct, drop = None, None
    elif abort is not None:
        drop_status = "not run (run aborted)"
        change_pct, drop = None, None
    elif baseline is None:
        drop_status = "no baseline"
        change_pct, drop = None, None
    else:
        change_pct, drop = drop_check(german, baseline)
        drop_status = "ALERT" if drop else "ok"
    if drop:
        alerts.append(drop)
    if limit_reached:
        alerts.append(
            f"REQUEST LIMIT REACHED ({api.requests_made} of {max_requests} live requests): the run stopped"
            " and is incomplete, it can never be a PASS"
        )

    def at(row):
        return {"max_count": int(row["max_count"]), "at": row["time"]} if row is not None else None

    return {
        "source": SOURCE_ORG,
        "mode": mode,
        "status": status,
        "outcome": outcome_label(status, grid_result, abort, count, mode),
        "status_reasons": reasons,
        "alerts": alerts,
        "aborted": abort is not None,
        "run_date": run_date,
        "started_at": started.isoformat(timespec="seconds"),
        "runtime_s": round(runtime_s, 1),
        "requests_made": api.requests_made,  # live requests only; cache hits are not requests
        "max_requests": max_requests,  # None: no limit (simulated)
        "request_limit_reached": limit_reached,
        "cache": "on" if cache is not None else "off",  # off: every query went to the API
        "cache_hits": cache.hits if cache is not None else 0,
        "cache_entries_loaded": cache.loaded if cache is not None else 0,
        "cache_entries_loaded_from": dict(cache.loaded_from) if cache is not None else {},
        "cache_hits_by_run": dict(cache.hits_by_run) if cache is not None else {},  # run folder -> answers
        "completion": _completion(grid_result, abort, early),
        "queries": len(log),
        "squares": {
            "complete": grid_result.count("complete"),
            "requery": grid_result.count("requery"),
            "split": grid_result.count("split"),
            "warning": grid_result.count("warning"),
            "error": grid_result.count("error"),
            "skipped": grid_result.squares_skipped,
            "not_queried": grid_result.squares_not_queried,
        },
        "coverage_warnings": grid_result.coverage_warnings,
        "open_squares": len(grid_result.open_squares()),  # listed in the open-squares CSV (files)
        "reference": {
            "zip": REFERENCE_ZIP,
            "city": REFERENCE_CITY,
            "radius_km": REFERENCE_RADIUS_KM,
            "expected_centre": list(REFERENCE_CENTRE),
            "actual_centre": list(centre) if centre else None,
            "centre_offset_km": reference_offset_km(reference) if reference else None,
            "start": at(start_row),
            "rechecks": [at(r) for r in recheck_rows],
            "end": at(end_row),
            "end_error": reference_end.get("error") if isinstance(reference_end, dict) else None,
            "end_centre_offset_km": end_offset,
            "countries_returned": _countries(reference.items) if reference else {},
        },
        "count_check": count,
        "early_stop": early,
        "unique_locations": unique,
        "unique_german_locations": german,
        "countries": df["country"].fillna("missing").value_counts().to_dict() if df is not None else {},
        # drop check vs the last full run: an alert, never changes status
        "drop_check": drop_status,
        "german_drop_alert": drop is not None,
        "baseline": (
            {k: baseline[k] for k in ("run_date", "started_at", "unique_german_locations")} | {"file": baseline["_file"]}
            if baseline
            else None
        ),
        "german_change_pct": change_pct,
        "geocoding_outliers": geocoding_outliers(log, GEOCODING_OUTLIER_KM),
        "site_type_counts": df["site_type"].value_counts().to_dict() if df is not None else {},
        "german_site_type_counts": (
            {t: int(german_df["site_type"].eq(t).sum()) for t in SITE_TYPES} if german_df is not None else {}
        ),
        "in_scope": int(df["in_scope"].sum()) if df is not None else None,  # German rows of the in-scope tiers
        "in_scope_by_tier": _tier_counts(german_df[german_df["in_scope"]]) if german_df is not None else {},
        "out_of_scope_by_reason": (
            german_df.loc[~german_df["in_scope"], "scope_reason"].fillna("none").value_counts().to_dict()
            if german_df is not None
            else {}
        ),
        "same_site_rows": int(german_df[SAME_SITE_COLUMN].notna().sum()) if german_df is not None else None,
        "needs_review_by_reason": enrich.reason_counts(df) if df is not None else {},
        "files": files,
        "renamed_files": [{"intended": a, "written": b} for a, b in renamed],
    }


def _outlier_square(o: dict) -> str:
    if o["square_km"] is None:
        return "reference query, no square"
    return f"{o['square_km']} km square lat {o['lat_min']}-{o['lat_max']} lng {o['lng_min']}-{o['lng_max']}"


def _reference_lines(summary: dict) -> list[str]:
    ref, count, early = summary["reference"], summary["count_check"], summary["early_stop"]

    def one(x):
        return f"{x['max_count']} (at {x['at']})" if x else "n/a"

    end = one(ref["end"])
    if ref["end_error"]:
        end = f"end query failed ({ref['end_error']})"
    lines = [
        f"Reference query ({ref['zip']} {ref['city']}, {ref['radius_km']} km):",
        f"  Actual centre {ref['actual_centre'][0]}, {ref['actual_centre'][1]} · {ref['centre_offset_km']} km from the expected"
        f" {ref['expected_centre'][0]}, {ref['expected_centre'][1]} (limit {REFERENCE_CENTRE_TOLERANCE_KM} km)"
        if ref["actual_centre"]
        else "  Actual centre n/a (no reference response)",
        f"  MaxCount at start {one(ref['start'])} · at end {end}",
    ]
    if ref["rechecks"]:
        lines.append("  Early-stop rechecks: " + ", ".join(one(r) for r in ref["rechecks"]))
    if count.get("evaluated"):
        span = (
            str(count["reference_low"])
            if count["reference_low"] == count["reference_high"]
            else f"{count['reference_low']}..{count['reference_high']}"
        )
        lines += [
            f"  N_circle {count['n_circle']} (collected IDs within {count['radius_km']} km of that centre) vs MaxCount {span}:"
            f" {count['verdict'].upper()}"
            + (f" (gap {count['gap']})" if count["gap"] else "")
            + (f" (excess {count['excess']})" if count["excess"] else ""),
            f"  Rows with Country != D: {count['abroad_inside']} inside the circle · {count['abroad_outside']} outside",
            f"  German total (Country == D): {count['german_total']}"
            + (" (= MaxCount)" if count["verdict"] == "match" and not count["abroad_inside"] else ""),
        ]
        if count["items_without_coordinates"]:
            lines.append(f"  ! {count['items_without_coordinates']} collected entries have no coordinates (not in N_circle)")
    else:
        lines.append("  N_circle: not evaluated (no reference response)")
    lines.append(f"  Address.Country of the {sum(ref['countries_returned'].values())} items it returned: "
                 f"{_join(ref['countries_returned']) or 'n/a'}")
    if early.get("requested"):
        if early.get("disabled_reason"):
            lines.append(f"Early stop: REQUESTED BUT DISABLED: {early['disabled_reason']}")
        else:
            state = "triggered" if early.get("triggered") else ("overshoot" if early.get("overshoot") else "not triggered")
            extra = f" · reference changed {early['target_changes']}" if early.get("target_changes") else ""
            extra += f" · off: {early['error']}" if early.get("error") else ""
            lines.append(f"Early stop: on, {state}, {early.get('rechecks', 0)} rechecks{extra}")
    else:
        lines.append("Early stop: off")
    probe = early.get("cap_probe")
    if probe:
        verdict = (
            "EQUAL: the cap is below what covers Germany, early stop would be disabled"
            if probe["equal"]
            else "different: the cap covers Germany"
        )
        lines.append(
            f"Cap probe München {CAP_PROBE_ZIP}: MaxCount {probe['max_counts'][0]} at {probe['radii_km'][0]} km,"
            f" {probe['max_counts'][1]} at {probe['radii_km'][1]} km ({verdict})"
        )
    return lines


def format_run_summary(summary: dict, max_warnings: int = 10) -> str:
    sq = summary["squares"]
    if summary["drop_check"] in ("ok", "ALERT"):
        b = summary["baseline"]
        baseline_line = (
            f"{summary['drop_check']}: {b['run_date']}: {b['unique_german_locations']} German -> now"
            f" {summary['unique_german_locations']} ({summary['german_change_pct']:+.1f}%, alert below -{MAX_GERMAN_DROP:.0%})"
        )
    elif summary["drop_check"] == "no baseline":
        baseline_line = "no baseline (no previous full run), drop check skipped"
    else:
        baseline_line = summary["drop_check"]
    german_types = summary["german_site_type_counts"]
    outliers = summary["geocoding_outliers"]
    lines = [
        f"GTÜ run summary ({summary['mode']}): {summary['outcome']}",
        "=" * 50,
        f"Run date:           {summary['run_date']}  (started {summary['started_at']}, {summary['runtime_s']} s)",
        f"Requests made:      {summary['requests_made']}"
        + (f" of max {summary['max_requests']}" if summary.get("max_requests") else "")
        + f"  (queries logged: {summary['queries']}, "
        + (
            f"cache: on, answered from the same-day cache: {summary['cache_hits']})"
            if summary.get("cache", "on") == "on"
            else "cache: off, every query sent live)"
        ),
        f"Completion:         {summary['completion']}",
        f"Squares:            complete {sq['complete']} · re-queried {sq['requery']} · split {sq['split']}"
        f" · warning {sq['warning']} · error {sq['error']} · skipped {sq['skipped']} · not queried {sq['not_queried']}",
        f"Open squares:       {summary['open_squares']}"
        + (f" (listed in {summary['files']['open_squares']})" if "open_squares" in summary["files"] else ""),
    ]
    if summary.get("cache_hits_by_run"):
        lines.append(
            "Cache answers by run folder: " + ", ".join(f"{run}: {n}" for run, n in summary["cache_hits_by_run"].items())
        )
    if summary["status_reasons"]:
        lines += ["FAIL because:"] + [f"  ! {r}" for r in summary["status_reasons"]]
    if summary["alerts"]:
        lines += ["ALERTS:"] + [f"  !! {a}" for a in summary["alerts"]]
    lines += [
        "",
        f"Collected:          {summary['unique_locations']} unique, {summary['unique_german_locations']} in Germany",
        f"Per Address.Country: {_join(summary['countries']) or 'n/a'}",
        f"German by site type: {_join(german_types) or 'n/a'}",
        f"In scope:           {summary['in_scope']} German rows ({_join(summary['in_scope_by_tier']) or 'n/a'})",
        f"Out of scope:       {_join(summary['out_of_scope_by_reason']) or 'none'}",
        f"Same-site rows:     {summary['same_site_rows']} German rows share a site with another row (info, not merged here)",
        "",
    ] + _reference_lines(summary) + [
        f"Last full run:      {baseline_line}",
        "",
        f"Site types (all):   {_join(summary['site_type_counts']) or 'n/a'}",
        f"Reasons:            {_join(summary['needs_review_by_reason']) or 'none'}",
        f"Geocoding:          {len(outliers)} queries with centre offset > {GEOCODING_OUTLIER_KM} km",
    ]
    lines += [
        f"  - PLZ {o['zip']} {o['city']}: {o['center_offset_km']} km, {_outlier_square(o)} (query {o['seq']})"
        for o in outliers
    ]
    warnings = summary["coverage_warnings"]
    if warnings:
        lines += ["", f"Coverage warnings ({len(warnings)}, all listed in the run JSON):"]
        lines += [
            f"  - {w['square_km']} km square lat {w['lat_min']}-{w['lat_max']} lng {w['lng_min']}-{w['lng_max']}"
            f" (PLZ {w['zip']}): {w['reason']}"
            for w in warnings[:max_warnings]
        ]
        if len(warnings) > max_warnings:
            lines.append(f"  ... {len(warnings) - max_warnings} more")
    lines += ["", "Files:"] + [f"  {k:<18}{v}" for k, v in summary["files"].items()]
    if summary["renamed_files"]:
        lines += ["", f"Existing files kept, new ones renamed ({len(summary['renamed_files'])}):"]
        lines += [f"  {r['intended']} -> {Path(r['written']).name}" for r in summary["renamed_files"]]
    return "\n".join(lines)


def render_summary_md(title: str, text: str) -> str:
    """run_summary.md: the printed summary under a heading. run_summary.json is the source of truth."""
    return f"# Run summary: {title}\n\n```text\n{text.rstrip()}\n```\n"


def run(
    mode: str = "simulated",
    confirm_full_run: bool = False,
    root: str | Path | None = None,
    run_date: str | None = None,
    centroids: pd.DataFrame | None = None,
    seed: int = 42,
    verbose: bool = True,
    sim_clusters: dict | None = None,
    use_cache: bool | None = None,
    early_stop: bool = False,
    requery_uncovered: bool = True,
    recheck_reference: bool | None = None,
    max_requests: int | None = None,
    started_at: datetime | None = None,
) -> RunResult:
    """Fetch the GTÜ finder with the adaptive grid, then parse, enrich and export.

    Every run starts with the reference query (Niederdorla, 500 km; see REFERENCE_*), which is always
    live, never cached. The run stops if its actual centre is more than 5 km off the expected one.
    Status rule C4: PASS needs (a) every grid square complete and (b) N_circle == the reference
    MaxCount (see run_status).

    Options:
      use_cache:          (default: on for mode="full", off otherwise) answer queries already made
                          today (raw.jsonl of earlier runs of the same date in the same area, output/ or
                          test_output/, plus this run) from disk instead of the API. Cache hits are not
                          requests and don't count against the limit. The reference and probe queries are
                          never cached. The summary says how many answers came from the cache and from
                          which run folders.
                          use_cache=False: every query, repeats included, goes to the API. The summary
                          shows "cache: off".
      started_at:         (default: now) start time that names the run folder (YYYY-MM-DD_HHMM_<mode>).
                          If only run_date is given, the folder date is run_date with the current time of day.
                          A same-minute clash gets the next free suffix (_2, _3, ...); a run never writes
                          into an existing folder.
      max_requests:       (default: 2,000 for mode="full", 5 for "sample"; ignored when simulated) safety
                          limit on live requests. If it is hit, the run stops, the summary says so, the
                          status is FAIL (never PASS) and RequestLimitReached is raised. The raw responses
                          up to that point are kept (see rebuild_from_raw).
      requery_uncovered:  (default on) grid.run_grid: re-query from the actual centre instead of
                          splitting when the count is complete but the radius doesn't cover the square.
      recheck_reference:  (default: on for mode="full") query the reference again at the end (1 request).
      early_stop:         (default off) stopping rule E: stop the grid once N_circle reaches the reference
                          MaxCount (see EarlyStop) and report PASS (early stop, count match) even with
                          squares still open; criterion (b) still applies. Needs a cap probe first
                          (2 extra requests: München 80686 at 455 and 460 km); equal counts disable it.
    A full run that isn't a PASS writes its files with a tag ("grid-incomplete" if squares are open) and
    lists the open squares in open_squares.csv (in the run folder).

    mode="simulated": no network, fake API with 1,200 locations spread like PLZ
                      centroids (realistic; must PASS). Stress test: pass
                      sim_clusters=grid.BIG_CITIES for 60 locations within 10 km of
                      each of 5 cities; all are found, but the run FAILs on purpose
                      with coverage warnings (more than 20 per smallest circle).
    mode="sample":    real API, 5 requests: reference query + cap probe (2) + 2 grid squares (testing
                      only). The probe runs in every sample, so its result is on record before a full
                      run; the early stop itself stays off unless early_stop=True.
    mode="full":      real API, whole grid; needs confirm_full_run=True.
    Loads the GeoNames centroids itself unless they are passed in.
    The summary is printed (verbose) and saved as run_summary.json and run_summary.md in the run folder.
    """
    if mode not in MODES:
        raise ValueError(f"mode must be one of {MODES}, got {mode!r}")
    if mode == "full" and not confirm_full_run:
        raise RuntimeError(FULL_RUN_NOT_CONFIRMED)

    if max_requests is not None and (isinstance(max_requests, bool) or not isinstance(max_requests, int) or max_requests < 1):
        raise ValueError(f"max_requests must be a positive integer, got {max_requests!r}")

    root = Path(root) if root is not None else PROJECT_ROOT
    stamp = started_at or datetime.now()
    if run_date is None:
        run_date = stamp.date().isoformat()
    elif started_at is None:
        stamp = datetime.combine(date.fromisoformat(run_date), stamp.time())
    started = stamp
    t0 = time.monotonic()
    if centroids is None:
        centroids = load_plz_centroids(ensure_geonames(paths.geonames_dir(root)))
    renamed: list = []  # (intended, written) for files that already existed

    if mode == "simulated":
        locations = simulated_locations(centroids, seed=seed, clusters=sim_clusters or {})
        api = FakeGtueApi(locations, centroids, seed=seed)
        raw_store = None  # fake responses are not worth keeping
        max_queries = None
        limit = None
    else:
        limit = max_requests or (SAMPLE_MAX_REQUESTS if mode == "sample" else FULL_MAX_REQUESTS)
        api = GtueClient(load_api_key(root), max_requests=limit)
        # sample: the reference query and the cap probe come first, the grid gets SAMPLE_GRID_SQUARES queries
        max_queries = SAMPLE_GRID_SQUARES if mode == "sample" else None
    run_folder = paths.new_run_dir(SLUG, mode, started, root)  # a new folder, never an existing one
    rp = run_paths(run_folder)
    if mode != "simulated":
        raw_store = RawStore(run_folder, root)
    if use_cache is None:
        use_cache = mode == "full"
    if recheck_reference is None:
        recheck_reference = mode == "full"
    cache = None
    if use_cache:
        cache = ResponseCache(root)
        cache.load_runs(paths.output_dir(SLUG, mode, root), run_date, exclude=run_folder)
    # The reference rechecks swallow StopRun (grid.reference_recheck, EarlyStop), so a limit hit there would
    # go unnoticed: every query function reports it here, and it always ends the run as an abort.
    limit_hit: list = []

    def watch_limit(fn):
        def watched(point, radius_km):
            try:
                return fn(point, radius_km)
            except RequestLimitReached as exc:
                limit_hit.append(exc)
                raise

        return watched

    query_fn = watch_limit(make_query_fn(api, raw_store, cache))
    live_fn = watch_limit(make_query_fn(api, raw_store))  # reference and probe answers must be fresh: never cached
    ref_point = QueryPoint(REFERENCE_ZIP, REFERENCE_CITY, *REFERENCE_CENTRE)

    grid_result = GridResult()
    reference = reference_end = None
    early = None
    early_info: dict = {"requested": early_stop, "disabled_reason": None, "cap_probe": None}
    abort = None
    try:
        reference = reference_check(live_fn, ref_point, REFERENCE_RADIUS_KM, grid_result)
        check_reference_centre(reference)
        if early_stop or mode == "sample":
            probe = run_cap_probe(live_fn, _query_point(centroids, CAP_PROBE_ZIP), grid_result)
            early_info["cap_probe"] = probe
            if early_stop and probe["equal"]:
                early_info["disabled_reason"] = (
                    f"cap probe: München MaxCount {probe['max_counts'][0]} at {probe['radii_km'][0]} km equals"
                    f" {probe['max_counts'][1]} at {probe['radii_km'][1]} km, so the server's radius cap has dropped"
                    f" below what covers Germany from {REFERENCE_CITY}; early stop disabled, the full grid runs"
                )
            elif early_stop:
                early = EarlyStop(live_fn, ref_point, reference)
        run_grid(
            query_fn,
            centroids,
            max_queries=max_queries,
            result=grid_result,
            requery_uncovered=requery_uncovered,
            stop_when=early,
        )
    except (StopRun, KeyboardInterrupt) as exc:
        abort = exc
    if early is not None:
        early_info.update(early.summary())

    if abort is None and recheck_reference:
        reference_end = reference_recheck(live_fn, ref_point, REFERENCE_RADIUS_KM, grid_result, status="reference_end")

    limit_reached = bool(limit_hit)
    if limit_reached and (abort is None or isinstance(abort, RequestLimitReached)):
        abort = RequestLimitReached(
            f"Request limit reached: {api.requests_made} live requests sent (max_requests={limit}). The run stopped"
            f" before it finished and is INCOMPLETE (FAIL, never PASS). The responses so far are in"
            f" {_relative(run_folder / paths.RAW_FILE, root)}; to continue, run again with a higher max_requests"
            " (or rebuild what was collected with rebuild_from_raw)."
        )

    query_log = pd.DataFrame(grid_result.log, columns=LOG_COLUMNS)
    log_path = safe_path(rp["query_log"], renamed)
    summary_path = safe_path(rp["summary"], renamed)
    query_log.to_csv(log_path, sep=";", encoding="utf-8-sig", index=False)
    files = {
        "run_folder": _relative(run_folder, root),
        "query_log": _relative(log_path, root),
        "run_summary": _relative(summary_path, root),
    }
    if raw_store is not None:
        files["raw_jsonl"] = _relative(raw_store.path, root)
    open_squares = grid_result.open_squares()
    if open_squares:
        open_path = safe_path(rp["open_squares"], renamed)
        pd.DataFrame(open_squares, columns=OPEN_SQUARE_COLUMNS).to_csv(
            open_path, sep=";", encoding="utf-8-sig", index=False
        )
        files["open_squares"] = _relative(open_path, root)

    df = None
    if abort is None:
        items = {"AddressItems": list(grid_result.items.values())}
        df = enrich_german(parse_response(items, scraped_at=started.isoformat(timespec="seconds")), centroids)
        # a full run that is not a PASS must not look like a complete crawl
        passed = _evaluate(mode, grid_result, abort, reference, reference_end)["status"] == "PASS"
        tag = None if passed or mode != "full" else (REBUILD_TAG if open_squares else UNVERIFIED_TAG)
        for key, (path, _rows) in export(df, run_folder, run_date, renamed, tag=tag).items():
            files[key] = _relative(path, root)

    baseline = _baseline(paths.output_dir(SLUG, mode, root), exclude=run_folder) if mode == "full" else None
    summary = _build_summary(
        mode, run_date, started, time.monotonic() - t0, api, grid_result, reference, df, abort, files, renamed, baseline,
        cache, reference_end, early_info, limit, limit_reached,
    )
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    text = format_run_summary(summary)
    rp["summary_md"].write_text(render_summary_md(run_folder.name, text), encoding="utf-8")
    if verbose:
        print(text)
    if abort is not None:
        raise abort
    return RunResult(summary, df, query_log)


# --------------------------------------------------------------------------
# Rebuild from raw (no network)
# --------------------------------------------------------------------------

REBUILD_TAG = "grid-incomplete"
UNVERIFIED_TAG = "count-mismatch"  # full run that is not a PASS although no square is open
AUTO_TAG = "auto"  # rebuild_from_raw: tag only if the source run did not pass (see its docstring)


def _source_passed(source: Path, n_unique: int, n_german: int) -> str | None:
    """The source run's outcome text if it passed and the rebuild finds its location counts, else None."""
    try:
        summary = json.loads((source / paths.RUN_SUMMARY_JSON).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    outcome = summary.get("outcome")
    if not (isinstance(outcome, str) and outcome.upper().startswith("PASS")):
        return None
    if summary.get("unique_locations") != n_unique or summary.get("unique_german_locations") != n_german:
        return None
    return outcome
@dataclass
class RebuildResult:
    summary: dict
    df: pd.DataFrame


def _read_raw_items(raw_jsonl: Path) -> tuple[dict, dict]:
    """All raw responses of a run in request order -> ({id: raw item}, counts). First seen wins,
    like the grid. Unreadable lines and entries without ProviderForeignKey are counted, not fatal."""
    items: dict = {}
    counts = {"files": 0, "unreadable": [], "items_seen": 0, "items_without_id": 0}
    for n, (name, data) in enumerate(iter_raw(raw_jsonl), start=1):
        try:
            entries = list(data["AddressItems"])
            if data.get("Exceptions") or data.get("InternalError"):
                raise ValueError("API reported an error")
        except (ValueError, KeyError, TypeError, AttributeError):
            counts["unreadable"].append(name or f"line {n}")
            continue
        counts["files"] += 1
        for item in entries:
            counts["items_seen"] += 1
            pfk = _clean(((item.get("BasicData") or {}).get("Identifiers") or {}).get("ProviderForeignKey"))
            if pfk is None:
                counts["items_without_id"] += 1
                continue
            items.setdefault(pfk, item)
    return items, counts


def rebuild_from_raw(
    source_run: str | Path,
    centroids: pd.DataFrame,
    root: str | Path | None = None,
    run_date: str | None = None,
    query_log: str | Path | None = None,
    tag: str | None = AUTO_TAG,
    verbose: bool = True,
    started_at: datetime | None = None,
) -> RebuildResult:
    """Parse, dedupe, classify and export the raw.jsonl of a run folder. No network, nothing overwritten.

    For a run that stopped before the grid was complete, or to re-classify a run that passed (a scope
    rule changed). The result goes to a new run folder of its own (mode "rebuild", next to its source
    run in output/) with a run_summary.json/.md that name the source run. It gets no copy of raw.jsonl.
    `source_run` is a run folder (a name is looked up in output/). scraped_at is the source run's start
    time (its run_summary.json), else the raw.jsonl's modification time. query_log (default: the source
    run's query_log.csv) adds the grid's square counts to the summary.

    `tag` (default: automatic): the exports carry `tag` in the file names. Automatic means no tag and
    outcome "PASS (rebuilt ...)" if the source run's run_summary.json says PASS and the rebuild finds
    exactly its unique and German location counts (nothing was lost, only the classification differs);
    otherwise REBUILD_TAG and "GRID INCOMPLETE", as for an unfinished run.
    """
    root = Path(root) if root is not None else PROJECT_ROOT
    source = Path(source_run)
    if not source.is_dir() and not source.is_absolute():
        source = paths.output_dir(SLUG, "full", root) / source
    if not source.is_dir():
        raise FileNotFoundError(f"Run folder not found: {source_run}")
    raw_jsonl = source / paths.RAW_FILE
    if not raw_jsonl.is_file():
        raise FileNotFoundError(f"No {paths.RAW_FILE} in {source}")
    run_date = run_date or source.name[:10]
    renamed: list = []

    items, counts = _read_raw_items(raw_jsonl)
    if not items:
        raise ValueError(f"No usable locations in {raw_jsonl}")
    scraped_at = datetime.fromtimestamp(raw_jsonl.stat().st_mtime).isoformat(timespec="seconds")
    try:
        scraped_at = str(json.loads((source / paths.RUN_SUMMARY_JSON).read_text(encoding="utf-8"))["started_at"])
    except (OSError, ValueError, KeyError):
        pass
    df = enrich_german(parse_response({"AddressItems": list(items.values())}, scraped_at), centroids)

    inherited = _source_passed(source, len(df), int(df["country"].eq(HOME_COUNTRY).sum()))
    if tag == AUTO_TAG:
        tag = None if inherited else REBUILD_TAG
    elif tag is None and not inherited:
        raise ValueError("tag=None (no tag) is only allowed for a source run that passed with the same counts")

    run_folder = paths.new_run_dir(SLUG, "rebuild", started_at or datetime.now(), root)
    files = {"run_folder": _relative(run_folder, root)} | {
        key: _relative(path, root)
        for key, (path, _rows) in export(df, run_folder, run_date, renamed, tag=tag).items()
    }

    # The first raw line is the reference query if its centre is where the reference should be (its
    # radius is not stored in the response; the reference radius is assumed). Older runs began with a
    # national check from elsewhere: then there is no circle to count in.
    first_max, first_centre, is_reference, first_name = None, None, False, None
    for first_name, first in iter_raw(raw_jsonl):
        try:
            first_max = int(first["Paging"]["MaxCount"])
            first_centre = (float(first["SearchData"]["Location"]["YCoord"]), float(first["SearchData"]["Location"]["XCoord"]))
            is_reference = float(haversine_km(*first_centre, *REFERENCE_CENTRE)) <= REFERENCE_CENTRE_TOLERANCE_KM
        except (ValueError, KeyError, TypeError):
            pass
        break
    count = (
        count_check(items, first_centre, [first_max]) if is_reference else {"evaluated": False, "problems": []}
    )

    squares = None
    query_log = Path(query_log) if query_log is not None else source / paths.QUERY_LOG_FILE
    if query_log.exists():
        log = pd.read_csv(query_log, sep=";", encoding="utf-8-sig")
        squares = {k: int(v) for k, v in log["status"].value_counts().items()}

    german = df[df["country"].eq(HOME_COUNTRY)]
    summary = {
        "source": SOURCE_ORG,
        "mode": "rebuild",
        "source_run": source.name,
        "grid_status": "complete (inherited from the source run)" if inherited else "incomplete",
        "outcome": f"PASS (rebuilt from {source.name}, re-classified, no network; source: {inherited})" if inherited else None,
        "label": tag,
        "run_date": run_date,
        "started_at": (started_at or datetime.now()).isoformat(timespec="seconds"),
        "source_raw": _relative(raw_jsonl, root),
        "raw_files_read": counts["files"],
        "raw_files_unreadable": counts["unreadable"],
        "items_seen": counts["items_seen"],
        "items_without_id": counts["items_without_id"],
        "unique_locations": len(df),
        "unique_german_locations": len(german),
        "first_response": {
            "file": first_name,
            "max_count": first_max,
            "centre": list(first_centre) if first_centre else None,
            "is_reference_query": is_reference,
        },
        "count_check": count,
        "query_log_statuses": squares,
        "countries": df["country"].fillna("missing").value_counts().to_dict(),
        "site_type_counts": df["site_type"].value_counts().to_dict(),
        "german_site_type_counts": {t: int(german["site_type"].eq(t).sum()) for t in SITE_TYPES},
        "in_scope": int(df["in_scope"].sum()),
        "in_scope_by_tier": _tier_counts(german[german["in_scope"]]),
        "out_of_scope_by_reason": german.loc[~german["in_scope"], "scope_reason"].fillna("none").value_counts().to_dict(),
        "same_site_rows": int(german[SAME_SITE_COLUMN].notna().sum()),
        "needs_review_count": int(df["needs_review"].sum()),
        "needs_review_by_reason": enrich.reason_counts(df),
        "files": files,
        "renamed_files": [{"intended": a, "written": b} for a, b in renamed],
    }
    summary_path = run_folder / paths.RUN_SUMMARY_JSON
    summary["files"]["run_summary"] = _relative(summary_path, root)
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    text = format_rebuild_summary(summary)
    (run_folder / paths.RUN_SUMMARY_MD).write_text(render_summary_md(run_folder.name, text), encoding="utf-8")
    if verbose:
        print(text)
    return RebuildResult(summary, df)


def _rebuild_count_lines(summary: dict) -> list[str]:
    first, count = summary["first_response"], summary["count_check"]
    if not count.get("evaluated"):
        why = "not evaluated: the first raw file is not a reference query (no circle to count in)"
        if first["max_count"] is not None:
            why += f" · its MaxCount {first['max_count']}"
        return [f"Count check:        {why}"]
    lines = [
        f"Reference MaxCount: {count['reference_low']} (first raw file) · N_circle {count['n_circle']}"
        f" (collected IDs within {count['radius_km']} km of its centre): {count['verdict'].upper()}"
        + (f" (gap {count['gap']})" if count["gap"] else "")
        + (f" (excess {count['excess']})" if count["excess"] else ""),
        f"Rows with Country != D: {count['abroad_inside']} inside the circle · {count['abroad_outside']} outside"
        f" · German total {count['german_total']}",
    ]
    return lines


def format_rebuild_summary(summary: dict) -> str:
    inherited = bool(summary.get("outcome"))
    lines = (
        [f"GTÜ run summary (rebuild): {summary['outcome']}", "=" * 50]
        if inherited
        else [
            "!" * 64,
            "GRID INCOMPLETE: rebuilt from raw responses of an unfinished run.",
            "Not a verified full crawl; the grid did not prove every square complete.",
            "!" * 64,
        ]
    ) + [
        f"GTÜ rebuild from raw ({summary['run_date']}, no network)",
        "=" * 50,
        f"Source run:         {summary['source_run']}  ({summary['raw_files_read']} raw responses read,"
        f" {len(summary['raw_files_unreadable'])} unreadable)",
        f"Entries seen:       {summary['items_seen']}  ({summary['items_without_id']} without ID)",
    ]
    if summary["query_log_statuses"]:
        lines.append(f"Query log:          {_join(summary['query_log_statuses'])}")
    lines += [
        "",
        f"Collected:          {summary['unique_locations']} unique, {summary['unique_german_locations']} in Germany",
        *_rebuild_count_lines(summary),
        f"Per Address.Country: {_join(summary['countries']) or 'n/a'}",
        "",
        f"Site types (all):   {_join(summary['site_type_counts']) or 'n/a'}",
        f"German by site type: {_join(summary['german_site_type_counts']) or 'n/a'}",
        f"In scope:           {summary['in_scope']} German rows ({_join(summary['in_scope_by_tier']) or 'n/a'})",
        f"Out of scope:       {_join(summary['out_of_scope_by_reason']) or 'none'}",
        f"Same-site rows:     {summary['same_site_rows']} German rows share a site with another row (info, not merged here)",
        f"Needs review:       {summary['needs_review_count']}",
        f"Reasons:            {_join(summary['needs_review_by_reason']) or 'none'}",
        "",
        "Files:",
    ]
    lines += [f"  {k:<18}{v}" for k, v in summary["files"].items()]
    if summary["renamed_files"]:
        lines += ["", f"Existing files kept, new ones renamed ({len(summary['renamed_files'])}):"]
        lines += [f"  {r['intended']} -> {Path(r['written']).name}" for r in summary["renamed_files"]]
    if inherited:
        lines += ["", "PASS is inherited: the source run passed and this rebuild found exactly its location counts."]
    else:
        lines += ["", "GRID INCOMPLETE: do not treat these files as the result of a complete run."]
    return "\n".join(lines)
