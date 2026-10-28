"""Shared export and run-history helpers (CSV/XLSX exports, no-overwrite file names, baseline run, drop check).

The GTÜ and DEKRA adapters each carry their own copy of these; new adapters (KÜS first) use this module.
"""

from __future__ import annotations

import json
import logging
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from core import paths

logger = logging.getLogger(__name__)

XLSX_TEXT_COLUMNS = [  # stored as text so Excel keeps leading zeros
    "source_id", "location_id", "zip", "location_sign", "aspro_id", "market_area_id", "federal_state_id", "company_id",
]
XLSX_MAX_COL_WIDTH = 60
EXPORT_KEYS = ("all_csv", "all_xlsx", "pruefstellen_csv", "pruefstellen_xlsx")
MAX_DROP = 0.10  # alert if in-scope rows drop more than this vs the last passing full run


def safe_path(path: str | Path, renamed: list | None = None) -> Path:
    """Never overwrite: if path exists, return name_2, name_3, ... and log the rename."""
    path = Path(path)
    if not path.exists():
        return path
    n = 2
    while (candidate := path.with_name(f"{path.stem}_{n}{path.suffix}")).exists():
        n += 1
    logger.warning("%s already exists, writing %s instead", path, candidate.name)
    if renamed is not None:
        renamed.append((str(path), str(candidate)))
    return candidate


def write_xlsx(frame: pd.DataFrame, path: Path, sheet_name: str) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(frame.columns, start=1):
            letter = get_column_letter(idx)
            if col in XLSX_TEXT_COLUMNS:
                for cell in ws[letter][1:]:
                    cell.number_format = "@"
                    if cell.value is not None:
                        cell.value = str(cell.value)
            width = max([len(str(col))] + [len(str(v)) for v in frame[col].dropna()])
            ws.column_dimensions[letter].width = min(width + 2, XLSX_MAX_COL_WIDTH)


def export(
    df: pd.DataFrame,
    out_dir: str | Path,
    slug: str,
    run_date: str | None = None,
    renamed: list | None = None,
    tag: str | None = None,
    only: tuple[str, ...] | list[str] | None = None,
) -> dict:
    """Write all rows and the in-scope rows to CSV (utf-8-sig, ';') and XLSX: `<slug>_all_<date>[_<tag>].csv`, ...

    only: a subset of EXPORT_KEYS (default all four). Existing files are never overwritten. Returns {key: (path, rows)}.
    """
    wanted = set(EXPORT_KEYS if only is None else only)
    unknown = wanted - set(EXPORT_KEYS)
    if unknown:
        raise ValueError(f"unknown export key(s) {sorted(unknown)}; choose from {EXPORT_KEYS}")
    run_date = run_date or date.today().isoformat()
    suffix = f"_{tag}" if tag else ""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    result = {}
    for key, frame in {"all": df, "pruefstellen": df[df["in_scope"]]}.items():
        if f"{key}_csv" in wanted:
            csv_path = safe_path(out_dir / f"{slug}_{key}_{run_date}{suffix}.csv", renamed)
            frame.to_csv(csv_path, sep=";", encoding="utf-8-sig", index=False)
            result[f"{key}_csv"] = (csv_path, len(frame))
        if f"{key}_xlsx" in wanted:
            xlsx_path = safe_path(out_dir / f"{slug}_{key}_{run_date}{suffix}.xlsx", renamed)
            write_xlsx(frame, xlsx_path, sheet_name=key)
            result[f"{key}_xlsx"] = (xlsx_path, len(frame))
    return result


def relative(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError:
        return str(path)


def baseline_run(output_dir: Path, exclude: Path | None = None) -> dict | None:
    """The newest earlier official run whose run_summary.json says mode full and PASS (with an in_scope count)."""
    if not output_dir.is_dir():
        return None
    for folder in sorted(output_dir.iterdir(), reverse=True):
        if folder == exclude or not (folder / paths.RUN_SUMMARY_JSON).exists():
            continue
        try:
            summary = json.loads((folder / paths.RUN_SUMMARY_JSON).read_text(encoding="utf-8"))
        except ValueError:
            continue
        if summary.get("mode") == "full" and summary.get("status") == "PASS" and summary.get("in_scope") is not None:
            return {"run_folder": folder.name, "run_date": summary.get("run_date"), "in_scope": summary["in_scope"]}
    return None


def drop_check(in_scope: int, baseline: dict | None, max_drop: float = MAX_DROP) -> tuple[str, bool]:
    """(text, alert): in-scope rows against the baseline run; alert when they fall by more than max_drop."""
    if not baseline:
        return "no baseline (no earlier passing full run)", False
    base = baseline["in_scope"]
    change = (in_scope - base) / base if base else 0.0
    alert = change < -max_drop
    return f"{'ALERT' if alert else 'ok'}: {base} -> {in_scope} in scope ({change:+.1%}, alert below -{max_drop:.0%})", alert


def counts(series: pd.Series) -> dict:
    return {str(k): int(v) for k, v in series.value_counts(dropna=True).items()}
